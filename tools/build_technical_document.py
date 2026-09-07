from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "bankinsight-technical-report"
OUT_MD = DOCS / "BankInsight_NL2SQL技术文档.md"
OUT_DOCX = DOCS / "BankInsight_NL2SQL技术文档.docx"

REPORT_TITLE = "BankInsight 金融指标智能问数系统技术文档"
REPORT_SUBTITLE = "面向第五届研究生金融科技创新大赛的技术路线、系统架构、实验设计与成果量化"
VERSION = "V1.0 · 证据快照 2026-08-27"

NAVY = "17365D"
BLUE = "2E74B5"
TEAL = "159E9C"
GOLD = "E3A62F"
INK = "243447"
MUTED = "5F6B7A"
LIGHT = "F4F6F9"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F6F5"
PALE_GOLD = "FFF5DD"
WHITE = "FFFFFF"
RED = "C0392B"
GREEN = "2E7D32"


def rgb(hex_color: str) -> tuple[int, int, int]:
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for ch in text:
        test = current + ch
        if draw.textlength(test, font=font) <= max_width or not current:
            current = test
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    fill: str,
    outline: str = BLUE,
    title_color: str = NAVY,
    body_color: str = INK,
    radius: int = 20,
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=rgb(fill), outline=rgb(outline), width=3)
    title_font = load_font(30, bold=True)
    body_font = load_font(22)
    draw.text((x1 + 24, y1 + 18), title, font=title_font, fill=rgb(title_color))
    y = y1 + 66
    for line in wrap_text(draw, body, body_font, x2 - x1 - 48):
        draw.text((x1 + 24, y), line, font=body_font, fill=rgb(body_color))
        y += 32


def draw_centered_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    fill: str,
    outline: str = BLUE,
    title_color: str = NAVY,
    body_color: str = INK,
    title_size: int = 30,
    body_size: int = 22,
    radius: int = 20,
) -> None:
    """Draw a box whose explicit text lines are centered in both directions."""
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=rgb(fill), outline=rgb(outline), width=3)

    title_font = load_font(title_size, bold=True)
    body_font = load_font(body_size)
    body_lines = body.splitlines() if body else []
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    title_height = title_bbox[3] - title_bbox[1]
    body_bboxes = [draw.textbbox((0, 0), line, font=body_font) for line in body_lines]
    body_heights = [bbox[3] - bbox[1] for bbox in body_bboxes]
    title_gap = 12 if body_lines else 0
    line_gap = 7
    block_height = title_height + title_gap + sum(body_heights) + line_gap * max(0, len(body_lines) - 1)
    current_y = y1 + ((y2 - y1) - block_height) / 2
    center_x = (x1 + x2) / 2

    title_width = title_bbox[2] - title_bbox[0]
    draw.text(
        (center_x - title_width / 2 - title_bbox[0], current_y - title_bbox[1]),
        title,
        font=title_font,
        fill=rgb(title_color),
    )
    current_y += title_height + title_gap

    for line, bbox, line_height in zip(body_lines, body_bboxes, body_heights):
        line_width = bbox[2] - bbox[0]
        draw.text(
            (center_x - line_width / 2 - bbox[0], current_y - bbox[1]),
            line,
            font=body_font,
            fill=rgb(body_color),
        )
        current_y += line_height + line_gap


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE, width: int = 5) -> None:
    draw.line([start, end], fill=rgb(color), width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    spread = 0.55
    p1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=rgb(color))


def save_canvas(image: Image.Image, name: str) -> Path:
    path = ASSETS / name
    image.save(path, "PNG", optimize=True)
    return path


def create_architecture_figure() -> Path:
    im = Image.new("RGB", (2000, 1200), rgb(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 42), "BankInsight 总体技术架构", font=load_font(52, True), fill=rgb(NAVY))
    d.text((70, 112), "规则优先、模型补位；统一经过语义、安全与结果校验", font=load_font(26), fill=rgb(MUTED))

    layers = [
        ("交互与集成层", 180, 145, PALE_BLUE),
        ("产品服务层", 350, 145, PALE_TEAL),
        ("NL2SQL 核心层", 525, 220, "EEF2FB"),
        ("可信执行层", 780, 155, PALE_GOLD),
        ("数据与状态层", 970, 175, LIGHT),
    ]
    label_x1, label_x2 = 60, 285
    for label, y, height, fill in layers:
        d.rounded_rectangle((55, y, 1945, y + height), radius=18, fill=rgb(fill), outline=rgb("D4DEE8"), width=2)
        label_font = load_font(27, True)
        label_bbox = d.textbbox((0, 0), label, font=label_font)
        label_width = label_bbox[2] - label_bbox[0]
        label_height = label_bbox[3] - label_bbox[1]
        d.text(
            (
                (label_x1 + label_x2 - label_width) / 2 - label_bbox[0],
                y + (height - label_height) / 2 - label_bbox[1],
            ),
            label,
            font=label_font,
            fill=rgb(NAVY),
        )

    draw_centered_box(d, (330, 198, 800, 307), "Web 工作台", "会话 · 图表 · 导出 · 分享", WHITE)
    draw_centered_box(d, (835, 198, 1305, 307), "OpenAPI 接入", "30 个操作 · 统一请求与错误模型", WHITE)
    draw_centered_box(d, (1340, 198, 1810, 307), "管理端", "用户 · 审计 · 告警 · 数据发布", WHITE)

    draw_centered_box(
        d, (330, 368, 800, 477), "ProductQueryService", "认证 · 授权 · 脱敏 · 审计编排", WHITE, TEAL, title_size=27, body_size=21
    )
    draw_centered_box(d, (835, 368, 1305, 477), "会话与结果服务", "历史 · 分享 · 导出 · 最终回答流", WHITE, TEAL, body_size=21)
    draw_centered_box(d, (1340, 368, 1810, 477), "数据发布服务", "预览 → 确认 → 事务发布", WHITE, TEAL, body_size=21)

    core = [
        (315, "上下文路由", "ContextRouter"),
        (625, "结构化规划", "Rule / LLM Planner"),
        (935, "确定性编译", "SQL Compiler"),
        (1245, "四级校验", "Plan · SQL · 对齐 · 结果"),
        (1555, "答案与图表", "Answerer / Chart"),
    ]
    for x, title, body in core:
        draw_centered_box(d, (x, 550, x + 275, 720), title, body, WHITE, BLUE, title_size=29, body_size=20)
    for i in range(len(core) - 1):
        arrow(d, (core[i][0] + 280, 635), (core[i + 1][0] - 10, 635), TEAL, 5)

    draw_centered_box(
        d, (340, 798, 970, 917), "权限域内执行", "组织与指标范围收缩 · 只读 DuckDB · 预检", WHITE, GOLD, body_size=21
    )
    draw_centered_box(d, (1070, 798, 1700, 917), "风险与审计", "动态脱敏 · 冻结 · 7 类告警 · 哈希链", WHITE, GOLD, body_size=21)
    arrow(d, (975, 858), (1058, 858), GOLD, 5)

    draw_centered_box(d, (315, 995, 790, 1128), "DuckDB", "13 机构 · 21 指标 · 132,698 行", WHITE, MUTED, body_size=20)
    draw_centered_box(d, (850, 995, 1325, 1128), "SQLite", "会话 · 待补全 · 审计\n告警 · 分享", WHITE, MUTED, body_size=20)
    draw_centered_box(d, (1385, 995, 1860, 1128), "外部 LLM（可选）", "仅处理规则未覆盖与复杂改写", WHITE, MUTED, body_size=20)
    return save_canvas(im, "01_system_architecture.png")


def create_pipeline_figure() -> Path:
    im = Image.new("RGB", (1800, 1160), rgb(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 45), "单次问数的可信执行链", font=load_font(48, True), fill=rgb(NAVY))
    d.text((70, 108), "任何规则或模型生成的 SQL 都不能绕过同一安全闸门", font=load_font(25), fill=rgb(MUTED))

    steps = [
        ("1", "身份与请求", "解析用户、调用系统、request_id"),
        ("2", "上下文决策", "自包含 / 需上下文 / 不确定"),
        ("3", "槽位补全", "时间、机构、指标、分析动作"),
        ("4", "规则规划", "高置信 QueryPlan；毫秒级主路径"),
        ("5", "模型补位", "仅在规则未覆盖或歧义时调用"),
        ("6", "确定性编译", "QueryPlan → DuckDB SQL"),
        ("7", "权限裁剪", "角色、指标白名单、机构域、脱敏"),
        ("8", "四级校验", "计划、SQL、语义对齐、结果义务"),
        ("9", "只读执行", "临时授权视图、超时与行数保护"),
        ("10", "结果校验", "空值、单位、排序、覆盖性"),
        ("11", "答案与图表", "表格、趋势、排名、可解释摘要"),
        ("12", "审计闭环", "哈希链、告警、冻结、可追踪导出"),
    ]
    left_x, right_x = 120, 930
    box_w, box_h = 700, 122
    for idx, (num, title, body) in enumerate(steps):
        col = 0 if idx < 6 else 1
        row = idx if idx < 6 else idx - 6
        x = left_x if col == 0 else right_x
        y = 185 + row * 150
        fill = PALE_BLUE if idx < 3 else (PALE_TEAL if idx < 6 else (PALE_GOLD if idx < 10 else LIGHT))
        d.rounded_rectangle((x, y, x + box_w, y + box_h), radius=18, fill=rgb(fill), outline=rgb(BLUE), width=3)
        d.ellipse((x + 18, y + 28, x + 84, y + 94), fill=rgb(NAVY))
        num_font = load_font(27, True)
        num_w = d.textlength(num, font=num_font)
        d.text((x + 51 - num_w / 2, y + 42), num, font=num_font, fill=rgb(WHITE))
        d.text((x + 108, y + 20), title, font=load_font(29, True), fill=rgb(NAVY))
        d.text((x + 108, y + 68), body, font=load_font(21), fill=rgb(INK))
        if row < 5:
            arrow(d, (x + 350, y + box_h), (x + 350, y + 145), TEAL if col == 0 else GOLD, 4)
    arrow(d, (left_x + box_w, 185 + 5 * 150 + box_h // 2), (right_x - 15, 185 + box_h // 2), BLUE, 5)
    return save_canvas(im, "02_trusted_query_pipeline.png")


def create_timeline_figure() -> Path:
    im = Image.new("RGB", (1800, 980), rgb(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 45), "技术攻关演进脉络", font=load_font(48, True), fill=rgb(NAVY))
    d.text((70, 108), "从可运行原型，迭代为可评测、可治理、可部署的金融问数产品", font=load_font(25), fill=rgb(MUTED))
    y_line = 505
    d.line((130, y_line, 1670, y_line), fill=rgb(BLUE), width=8)
    events = [
        (170, "07-21", "规则优先原型", "语义目录、QueryPlan、SQL 编译与 DuckDB"),
        (420, "07-22", "多轮基线暴露", "50 题仅 26% 全正确；定位上下文污染"),
        (685, "08-17", "产品化架构", "会话、导出、分享、OpenAPI、Web 工作台"),
        (950, "08-24", "细粒度安全", "角色/机构/指标域、脱敏、审计哈希链、告警"),
        (1215, "08-25", "上下文重构", "三态路由、PendingQuery、结果覆盖校验"),
        (1480, "08-26—27", "交付闭环", "273 题 95.97%；Docker、HTTPS、增量发布"),
    ]
    for idx, (x, date, title, body) in enumerate(events):
        above = idx % 2 == 0
        d.ellipse((x - 22, y_line - 22, x + 22, y_line + 22), fill=rgb(TEAL), outline=rgb(WHITE), width=5)
        box_y1, box_y2 = (210, 430) if above else (580, 820)
        d.line((x, y_line - 22 if above else y_line + 22, x, box_y2 if above else box_y1), fill=rgb("9FB9D3"), width=3)
        box_x1 = max(70, x - 120)
        box_x2 = min(1730, x + 220)
        d.rounded_rectangle((box_x1, box_y1, box_x2, box_y2), radius=18, fill=rgb(PALE_BLUE if above else PALE_TEAL), outline=rgb(BLUE), width=2)
        d.text((box_x1 + 20, box_y1 + 18), date, font=load_font(25, True), fill=rgb(TEAL))
        d.text((box_x1 + 20, box_y1 + 57), title, font=load_font(28, True), fill=rgb(NAVY))
        y = box_y1 + 105
        for line in wrap_text(d, body, load_font(21), box_x2 - box_x1 - 40):
            d.text((box_x1 + 20, y), line, font=load_font(21), fill=rgb(INK))
            y += 31
    d.text((120, 890), "关键方法论：先用错误分类明确语义边界，再以统一校验链保证新增能力不会破坏安全基线。", font=load_font(25, True), fill=rgb(NAVY))
    return save_canvas(im, "03_evolution_timeline.png")


def create_metrics_figure() -> Path:
    im = Image.new("RGB", (1800, 1050), rgb(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 45), "评测结果与工程质量快照", font=load_font(48, True), fill=rgb(NAVY))
    d.text((70, 108), "准确率采用人工复核口径；工程测试反映 2026-08-27 当前工作树", font=load_font(25), fill=rgb(MUTED))

    charts = [
        ("200 题单轮复核", 96.00, "192 / 200 正确", BLUE),
        ("73 题多轮复核", 95.89, "70 / 73 正确", TEAL),
        ("273 题综合复核", 95.97, "262 / 273 正确", NAVY),
        ("当前后端测试", 99.58, "239 / 240 通过*", GOLD),
    ]
    x0, chart_w, gap = 100, 360, 50
    base_y, max_h = 790, 500
    for idx, (label, value, detail, color) in enumerate(charts):
        x = x0 + idx * (chart_w + gap)
        d.rounded_rectangle((x, 230, x + chart_w, 890), radius=18, fill=rgb(LIGHT), outline=rgb("D7E0E8"), width=2)
        bar_x1, bar_x2 = x + 95, x + chart_w - 95
        bar_h = int(max_h * value / 100)
        d.rounded_rectangle((bar_x1, base_y - bar_h, bar_x2, base_y), radius=18, fill=rgb(color))
        value_text = f"{value:.2f}%"
        value_font = load_font(38, True)
        value_w = d.textlength(value_text, font=value_font)
        d.text((x + (chart_w - value_w) / 2, 160), value_text, font=value_font, fill=rgb(color))
        label_font = load_font(26, True)
        label_w = d.textlength(label, font=label_font)
        d.text((x + (chart_w - label_w) / 2, 815), label, font=label_font, fill=rgb(NAVY))
        detail_font = load_font(21)
        detail_w = d.textlength(detail, font=detail_font)
        d.text((x + (chart_w - detail_w) / 2, 855), detail, font=detail_font, fill=rgb(MUTED))
    d.text((100, 945), "* 唯一失败项为固定断言 132,678 行；当前增量发布后数据库为 132,698 行，属于测试夹具未同步，不是查询逻辑失败。", font=load_font(22), fill=rgb(RED))
    return save_canvas(im, "04_experiment_metrics.png")


def create_security_figure() -> Path:
    im = Image.new("RGB", (2000, 1200), rgb(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 42), "纵深安全与数据治理", font=load_font(52, True), fill=rgb(NAVY))
    d.text((70, 112), "安全控制覆盖请求前、生成中、执行时、结果后四个阶段", font=load_font(26), fill=rgb(MUTED))

    rings = [
        (880, "审计与处置", "哈希链 · 7 类告警 · 自动冻结 · 管理闭环", "D9EAF7"),
        (700, "结果治理", "S3 动态脱敏 · 导出阈值 · 分享限流", "D7F0EB"),
        (520, "可信执行", "SQLGuard · 临时授权视图 · 只读 / 超时 / 行数", "FFF0C8"),
        (340, "最小权限", "6 类业务角色 · 指标白名单 · 机构域", "F4DADA"),
    ]
    cx, cy = 1000, 635
    for diameter, title, body, fill in rings:
        x1, y1 = cx - diameter // 2, cy - diameter // 2
        x2, y2 = cx + diameter // 2, cy + diameter // 2
        d.ellipse((x1, y1, x2, y2), fill=rgb(fill), outline=rgb(WHITE), width=8)
        title_font = load_font(26, True)
        body_font = load_font(19)
        title_bbox = d.textbbox((0, 0), title, font=title_font)
        body_bbox = d.textbbox((0, 0), body, font=body_font)
        title_w = title_bbox[2] - title_bbox[0]
        body_w = body_bbox[2] - body_bbox[0]
        title_y = y1 + 20
        body_y = y1 + 58
        d.text((cx - title_w / 2 - title_bbox[0], title_y - title_bbox[1]), title, font=title_font, fill=rgb(NAVY))
        d.text((cx - body_w / 2 - body_bbox[0], body_y - body_bbox[1]), body, font=body_font, fill=rgb(INK))

    draw_centered_box(
        d,
        (835, 570, 1165, 755),
        "受控查询核心",
        "先授权，后执行；\n规则与 LLM 同闸门",
        NAVY,
        NAVY,
        title_color=WHITE,
        body_color=WHITE,
        title_size=31,
        body_size=22,
        radius=28,
    )

    note = "边界声明：当前数据粒度为“机构—指标—日期—数值”，不等同于客户级行权限；SQLite 哈希链亦不等同于监管级 WORM。"
    d.rounded_rectangle((90, 1100, 1910, 1170), radius=18, fill=rgb("FDF1EF"), outline=rgb("F2C6C0"), width=2)
    note_font = load_font(22, True)
    note_bbox = d.textbbox((0, 0), note, font=note_font)
    note_w = note_bbox[2] - note_bbox[0]
    note_h = note_bbox[3] - note_bbox[1]
    d.text(
        (cx - note_w / 2 - note_bbox[0], 1135 - note_h / 2 - note_bbox[1]),
        note,
        font=note_font,
        fill=rgb(RED),
    )
    return save_canvas(im, "05_security_layers.png")


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = "D6DEE6", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_layout(table, widths: Sequence[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)
    for ridx, row in enumerate(table.rows):
        prevent_row_split(row)
        if ridx == 0:
            repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, latin: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*rgb(color))


def style_paragraph(paragraph, *, size=11, bold=False, color=INK, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, before=0, line=1.333) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.widow_control = True
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None, style: str | None = None, keep=False) -> object:
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 11, True, NAVY)
        r2 = p.add_run(text[len(bold_lead) :])
        set_run_font(r2, 11, False, INK)
    else:
        r = p.add_run(text)
        set_run_font(r, 11, False, INK)
    style_paragraph(p)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc: Document, text: str, level: int = 0) -> object:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, 10.5, False, INK)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.widow_control = True
    return p


def add_numbered(doc: Document, text: str, level: int = 0) -> object:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, 10.5, False, INK)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.widow_control = True
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> object:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def add_callout(doc: Document, title: str, body: str, fill: str = PALE_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    prevent_row_split(table.rows[0])
    cell.width = Inches(6.45)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=190, bottom=130, end=190)
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    set_run_font(r, 10.5, True, accent)
    r2 = p.add_run(body)
    set_run_font(r2, 10.5, False, INK)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], LIGHT)
        p = header_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_run_font(r, 9.2, True, NAVY)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
    for data_row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(data_row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, 8.8, False, INK)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
    set_repeat_table_layout(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def new_numbering_instance(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def add_numbered_list(doc: Document, items: Sequence[str]) -> None:
    num_id = new_numbering_instance(doc)
    for text in items:
        p = doc.add_paragraph()
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])
        p_pr.append(num_pr)
        r = p.add_run(text)
        set_run_font(r, 10.5, False, INK)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.widow_control = True


def add_picture(doc: Document, path: Path, caption: str, alt: str, width: float = 6.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if doc.inline_shapes:
        doc.inline_shapes[-1]._inline.docPr.set("descr", alt)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(caption)
    set_run_font(cr, 9, False, MUTED)
    cp.paragraph_format.space_after = Pt(10)
    cp.paragraph_format.keep_together = True


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(*rgb(INK))
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 12),
        ("Subtitle", 15, MUTED, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, NAVY, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor(*rgb(color))
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Caption Custom" not in styles:
        caption = styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Calibri"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        caption.font.size = Pt(9)
        caption.font.color.rgb = RGBColor(*rgb(MUTED))

    for list_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("BANKINSIGHT  /  TECHNICAL REPORT")
    set_run_font(r, 8.5, True, BLUE)
    hp.add_run("\t")
    rr = hp.add_run(VERSION)
    set_run_font(rr, 8.5, False, MUTED)
    hp.paragraph_format.space_after = Pt(0)
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5))

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("BankInsight 技术文档  ·  ")
    set_run_font(fr, 8.5, False, MUTED)
    add_field(fp, "PAGE")
    fp.paragraph_format.space_after = Pt(0)

    doc.core_properties.title = REPORT_TITLE
    doc.core_properties.subject = "金融指标 NL2SQL 技术路线、系统架构、实验设计与成果量化"
    doc.core_properties.author = "BankInsight 项目团队"
    doc.core_properties.keywords = "NL2SQL, 金融科技, DuckDB, 语义解析, 数据安全, OpenAPI"
    return doc


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("FINTECH · NL2SQL · TRUSTED ANALYTICS")
    set_run_font(r, 10, True, TEAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(REPORT_TITLE)
    set_run_font(r, 30, True, NAVY)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_together = True

    p = doc.add_paragraph()
    r = p.add_run(REPORT_SUBTITLE)
    set_run_font(r, 15, False, MUTED)
    p.paragraph_format.space_after = Pt(30)
    p.paragraph_format.line_spacing = 1.25

    meta = doc.add_paragraph()
    mr = meta.add_run("BankInsight 项目团队  ·  " + VERSION)
    set_run_font(mr, 9.5, True, MUTED)
    meta.paragraph_format.space_after = Pt(18)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    metrics = [("95.97%", "273 题人工复核准确率"), ("98.53%", "规则主路径占比"), ("132,698", "当前指标事实行数")]
    for idx, (value, label) in enumerate(metrics):
        cell = table.cell(0, idx)
        set_cell_shading(cell, PALE_BLUE if idx != 1 else PALE_TEAL)
        set_cell_margins(cell, top=170, start=110, bottom=170, end=110)
        cell.width = Inches(2.1)
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rv = p1.add_run(value)
        set_run_font(rv, 20, True, NAVY if idx != 1 else TEAL)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rl = p2.add_run(label)
        set_run_font(rl, 9, False, MUTED)
    set_table_borders(table, "FFFFFF", "0")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    r = p.add_run("核心主张")
    set_run_font(r, 12, True, BLUE)
    p.paragraph_format.space_after = Pt(8)
    claims = [
        "以结构化 QueryPlan 作为自然语言与 SQL 之间的可信中间层；",
        "以规则优先获得稳定性与低延迟，以 LLM 补齐长尾语义；",
        "以统一权限、SQL、语义与结果校验链约束所有执行路径；",
        "以可审计、可发布、可容器化的产品工程闭环承载参赛成果。",
    ]
    for claim in claims:
        add_bullet(doc, claim)

    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document) -> None:
    add_heading(doc, "文档导读", 1)
    add_paragraph(doc, "本文以当前代码、当前数据、仓库内评测记录及同项目历史会话为证据基础，系统还原从原型到交付的技术攻关过程。为避免把历史阶段结果误写成最终结论，全文将“当前可复现”“历史里程碑”“人工复核”“边界声明”分开标注。")
    add_callout(doc, "证据口径", "当前状态以 2026-08-27 工作树为准；历史数据用于说明演进，不替代当前复测。项目与会话中未发现官方获奖名次或证书，因此文档仅报告可验证的技术成果，不虚构赛事排名。", PALE_GOLD, GOLD)
    chapters = [
        ("1", "项目概述与问题定义", "业务目标、约束、成功标准"),
        ("2", "数据与金融语义底座", "数据规模、指标口径、衍生规则"),
        ("3", "总体技术路线", "规则优先、模型补位、统一可信闸门"),
        ("4", "模型架构与关键技术", "上下文、规划、编译、校验、执行、回答"),
        ("5", "产品化、接口与前端", "OpenAPI、会话、导出、分享、管理端"),
        ("6", "安全治理与部署", "最小权限、审计告警、Docker、HTTPS"),
        ("7", "实验设计与量化结果", "单轮、多轮、综合评测、工程验证"),
        ("8", "技术攻关脉络", "问题暴露、根因定位、迭代闭环"),
        ("9", "创新点、边界与后续路线", "可迁移价值与工程风险"),
        ("附录", "模块、角色、API 与指标口径", "审阅和复现索引"),
    ]
    add_table(doc, ["章节", "主题", "阅读重点"], chapters, [0.7, 2.0, 3.6])
    doc.add_page_break()


def build_docx(figs: dict[str, Path]) -> None:
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "执行摘要", 1)
    add_paragraph(doc, "BankInsight 是面向银行经营指标查询场景的垂直 NL2SQL 系统。它不把大模型直接生成 SQL 作为唯一主链，而是以金融语义目录和结构化 QueryPlan 为核心：高置信问题由规则规划器直接解析，复杂或未覆盖表达再由 LLM 补位；两条路径最终统一进入权限裁剪、SQL 安全、语义对齐和结果覆盖校验。该路线同时解决了比赛题目中的准确性、响应速度、可解释性，以及金融场景中的最小权限与审计要求。")
    add_paragraph(doc, "截至证据快照，系统覆盖 13 家机构、21 项指标、487 个日期、132,698 条指标事实；提供 29 条 API 路径、30 个 OpenAPI 操作和 6 类业务角色。人工复核的 273 题综合集达到 262 题正确、11 题口径不一致或未匹配、0 运行错误，准确率 95.97%；其中规则主路径覆盖 269 题，占 98.53%。")
    add_picture(doc, figs["metrics"], "图 1  当前人工复核结果与工程质量快照", "四组柱形图，分别展示200题单轮、73题多轮、273题综合以及当前后端测试通过率。")
    add_callout(doc, "结论先行", "项目的核心竞争力不是“能生成 SQL”，而是把语义理解、确定性编译、权限隔离、结果义务和产品交付组合成一个可验证闭环。", PALE_TEAL, TEAL)

    add_heading(doc, "1 项目概述与问题定义", 1)
    add_heading(doc, "1.1 业务目标", 2)
    add_paragraph(doc, "比赛数据由银行机构、指标、日期和数值构成。用户问题表面上是自然语言转 SQL，实际同时包含金融术语识别、时间归一、组织范围、统计动作、排名方向、单位换算、衍生指标和权限边界。系统的目标是让业务人员无需编写 SQL，即可完成查询、对比、排名、趋势、结构、环比/同比和勾稽关系分析，并获得表格、图表与自然语言结论。")
    add_heading(doc, "1.2 核心挑战", 2)
    challenges = [
        ("语义歧义", "“最高/最低”“全省/全行”“今年2月1日”等表达必须映射为明确且可审计的结构化槽位。"),
        ("多轮依赖", "省略指标、机构或日期时，既要继承正确上下文，又不能让上一轮成功或失败状态污染当前问题。"),
        ("精确计算", "排名方向、金额与比例单位、日均/求和、环比/同比、勾稽等必须确定性执行。"),
        ("金融安全", "不同岗位只能查询被授权的机构与指标；脱敏、导出、分享和审计需遵守同一权限语义。"),
        ("工程交付", "能力必须通过 API、Web 工作台、容器化、HTTPS、数据增量发布和运维审计真正可用。"),
    ]
    add_table(doc, ["挑战", "技术含义"], challenges, [1.35, 5.0])
    add_heading(doc, "1.3 设计原则与成功标准", 2)
    for text in [
        "可信优先：任何 SQL，无论来自规则还是 LLM，都必须经过统一校验与授权执行。",
        "确定性优先：金额、排名、比较、派生计算尽量交给规则、编译器和数据库，而非让模型口算。",
        "渐进增强：先以小而稳定的有限状态机解决当前业务，再按真实复杂度引入更重的编排框架。",
        "证据驱动：以人工复核准确率、错误率、路由占比、延迟分位数、测试结果和部署可达性量化成果。",
        "边界清晰：不把容器化等同于国产化认证，不把哈希链等同于监管级 WORM，不把机构级数据域表述为客户级行权限。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "2 数据与金融语义底座", 1)
    add_heading(doc, "2.1 当前数据规模", 2)
    rows = [
        ("机构数", "13", "organization_id / organization_name"),
        ("指标数", "21", "覆盖资产、负债、收入、风险、资本等主题"),
        ("衍生规则", "10", "组合、排名、阈值、增速与勾稽等"),
        ("指标事实", "132,698 行", "增量发布后较历史基线增加 20 行"),
        ("日期范围", "2024-12-31—2026-06-30", "487 个不同日期"),
        ("主键", "日期+机构+指标", "当前重复主键为 0"),
    ]
    add_table(doc, ["维度", "当前值", "说明"], rows, [1.35, 1.8, 3.2])
    add_paragraph(doc, "数据层采用 DuckDB，核心事实表字段为 DATE、organization_id、metric_id 和 DOUBLE 数值。完整 Excel 可重建数据库；当前同时支持管理端增量 Excel 预览与事务发布。增量文件可以只包含“指标数据表”并覆盖部分已有指标，但不能凭文件新增指标定义。")
    add_callout(doc, "数据治理提醒", "增量发布会改变数据库行数与日期覆盖；若随后从未同步的原始 Excel 全量重建，可能覆盖增量数据。因此生产流程应把“发布文件归档—源文件同步—重建校验”作为一个原子化运维规范。", PALE_GOLD, GOLD)

    add_heading(doc, "2.2 语义目录与口径优先级", 2)
    add_paragraph(doc, "SemanticCatalog 统一管理指标编码、别名、单位、方向属性、组织别名和派生规则。规划器不直接依赖表字段猜测含义，而是先把用户语言绑定到目录实体，再生成 QueryPlan。对“经营情况好/差”等口径，系统以官方衍生维度说明为最高优先级：多数指标越高越好，ZB012、ZB013、ZB017 越低越好；13 家机构中 RANK 1—3 为好、4—9 为中性、10—13 为差。")
    add_paragraph(doc, "仓库测试材料中，第 91、93、151、152、191、192 题的参考答案与官方衍生规则存在冲突：部分答案采用“优于平均值”或漏掉资本/逾期类指标。当前实现坚持官方规则而非针对冲突答案定向调参，并在人工复核中将此类问题标为“口径冲突”，这是防止比赛集过拟合的重要措施。")

    add_heading(doc, "2.3 QueryPlan：可信中间表示", 2)
    add_paragraph(doc, "QueryPlan 是系统的核心模型接口。它把自然语言拆为指标集合、机构集合、时间范围、查询形态、排序方向、Top/Bottom N、阈值、聚合方式、单位与派生计算等可验证字段。规划与执行分离后，系统能够在 SQL 生成前发现缺槽、越权和语义矛盾，也能用同一编译器复用规则和 LLM 结果。")
    query_types = [
        ("点查/多指标", "指定机构、日期、一个或多个指标"),
        ("排名/Top-Bottom", "按指标方向排序，支持全省与局部机构"),
        ("阈值/计数", "大于、小于、区间、满足条件机构数"),
        ("趋势/区间", "期间值、日均、求和、变化量"),
        ("同比/环比", "基期定位、差额、增幅与单位处理"),
        ("结构/占比", "分子分母指标组合与百分比格式"),
        ("勾稽/对账", "派生公式、左右项差异与覆盖说明"),
        ("画像/综合", "机构多指标概览与经营好坏分类"),
    ]
    add_table(doc, ["查询形态", "QueryPlan 表达"], query_types, [1.8, 4.55])

    add_heading(doc, "3 总体技术路线", 1)
    add_picture(doc, figs["architecture"], "图 2  BankInsight 总体技术架构", "五层系统架构图，从交互与集成层到数据与状态层，突出NL2SQL核心与可信执行层。")
    add_heading(doc, "3.1 规则优先、模型补位", 2)
    add_paragraph(doc, "高频金融问数具有强结构性，若全部交给 LLM，会引入随机性、长尾延迟和越权风险。因此系统优先使用 RulePlanner 解析高置信表达，并由 SQLCompiler 确定性生成 DuckDB SQL；规则未覆盖、表达歧义或复杂改写才进入 LLMPlanner。最终回答同样拆分为“数据事实”和“文字分析”：事实值由数据库结果决定，模型只能在受控结果摘要上组织语言。")
    add_paragraph(doc, "在 273 题人工复核集中，规则路径覆盖 269 题、LLM 路径 4 题，规则占比 98.53%。这使整体延迟中位数维持约 0.07 秒；少数 LLM 请求造成 20 秒级尾延迟，但不会拖慢绝大多数确定性查询。")
    add_heading(doc, "3.2 全链路可信闸门", 2)
    add_picture(doc, figs["pipeline"], "图 3  单次问数的十二步可信执行链", "十二步流程图，说明身份、上下文、规划、编译、授权、校验、执行、回答和审计的顺序。")
    add_paragraph(doc, "系统把安全和正确性放在 SQL 执行之前：身份与角色决定可见组织/指标；PlanValidator 检查结构化计划；SQLGuard 基于 sqlglot 解析 SQL 并限制语句形态、表和字段；AlignmentValidator 检查 SQL 与计划的指标、机构、日期和聚合是否一致；DuckDBExecutor 再把查询限制在临时授权视图中。执行完成后，ResultValidator 和 AnswerCoverageValidator 检查空结果、排序、单位、排名、差值、阈值、计数和勾稽义务是否被答案完整覆盖。")

    add_heading(doc, "4 模型架构与关键技术", 1)
    add_heading(doc, "4.1 上下文路由：先判断是否需要历史", 2)
    add_paragraph(doc, "早期实现只要识别失败就盲目继承上一轮，导致 last_metrics、last_date 或 last_orgs 为空时回退到错误默认值。重构后，ContextRouter 在规划前做独立三态判断：SELF_CONTAINED 表示问题可单独求解；CONTEXT_REQUIRED 表示存在指代、省略或续问；UNCERTAIN 表示需交给 LLM 改写或追问。只有被判定为上下文依赖的问题才读取历史槽位，从根本上减少“无关继承”。")
    add_table(doc, ["决策状态", "处理方式", "风险控制"], [
        ("SELF_CONTAINED", "直接解析当前问题", "禁止无条件继承上一轮槽位"),
        ("CONTEXT_REQUIRED", "显式槽位 > PendingQuery > 上次成功", "继承范围可解释、可审计"),
        ("UNCERTAIN", "LLM 改写或发起澄清", "不直接生成并执行模糊 SQL"),
    ], [1.4, 2.35, 2.6])

    add_heading(doc, "4.2 PendingQuery：失败状态与成功记忆分离", 2)
    add_paragraph(doc, "对缺机构、缺日期或缺指标的问题，系统不把失败请求写入成功历史，而是生成 PendingQuery，保存已确认槽位、缺失槽位和原问题。后续补充按“当前显式槽位 > 待补全状态 > 最近成功 > LLM 推断”合并。PendingQuery 按 user_id+session_id 持久化到 SQLite，后端重启后仍可继续补全。历史验证中，用户先询问缺机构的净利润问题，服务重启后再补充“J市”，系统仍返回 175.6 及 +25.34 的正确结果。")
    add_callout(doc, "状态机选择", "当前流程是有限且可枚举的“解析—补全—校验—执行—回答”状态机。项目评估后未引入 LangGraph；只有当出现人工审批、多智能体、长流程中断恢复或复杂分支编排时，才值得承担额外框架复杂度。", PALE_BLUE, BLUE)

    add_heading(doc, "4.3 规则规划与 LLM 规划", 2)
    add_paragraph(doc, "RulePlanner 覆盖点查、排名、多指标、省均、阈值、排名阈值、成员过滤、期间值、日均、求和、结构占比、环比/同比、勾稽、趋势和画像等高频形态。日期解析器支持“今年2月1日”“25年2月1日”等表达，组织解析覆盖“全省、全行、全市、13个市”等口径。")
    add_paragraph(doc, "LLMPlanner 只接收最小必要上下文：当前问题、最近成功问题与简短回答、结构化槽位和必要结果摘要；不传递完整 SQL、完整日志或失败执行内容，降低提示污染与敏感信息泄露。LLM 输出仍需转为 QueryPlan，并通过与规则路径相同的校验链。")

    add_heading(doc, "4.4 确定性 SQL 编译", 2)
    add_paragraph(doc, "SQLCompiler 根据 QueryPlan 选择固定模板、CTE、窗口函数和派生表达式。排名采用数据库 RANK 语义，方向由指标目录决定；比例统一在编译或答案格式化阶段转换，避免模型把 0.0242% 误写成 2.42% 或 242%；同比/环比显式定位基期并返回基值、现值和变化量。确定性编译使 SQL 可复现、可测试，也便于 SQLGuard 做结构化比对。")

    add_heading(doc, "4.5 四级验证与答案覆盖", 2)
    validators = [
        ("PlanValidator", "槽位完整性、查询形态、阈值/排名参数、时间与组织约束"),
        ("SQLGuard", "仅允许只读查询；限制表、字段、函数、子查询和危险语句；sqlglot AST 解析"),
        ("AlignmentValidator", "QueryPlan 与 SQL 中的指标、机构、日期、聚合、排序一致"),
        ("ResultValidator", "空结果、列结构、数值类型、排序与单位合理性"),
        ("AnswerCoverageValidator", "排名、差值、比较、阈值、计数、勾稽等问题义务必须出现在答案中"),
    ]
    add_table(doc, ["组件", "核心职责"], validators, [1.8, 4.55])
    add_paragraph(doc, "答案覆盖校验是本项目相较普通 NL2SQL 的重要增强。很多 SQL 实际执行成功，但回答只报一个数，遗漏“比谁高多少”“排名第几”“有几家满足条件”等用户真正问到的义务。覆盖校验把这些义务结构化，并要求最终回答逐项给出，避免“SQL 对、答案不完整”。")

    add_heading(doc, "4.6 结果生成、图表与解释", 2)
    add_paragraph(doc, "结果层根据查询形态选择点值、明细表、排名条形图或时间趋势图，并生成简明结论。对复杂分析，FinalAnswerService 可在受控结果上生成解释；流式回答与数据查询解耦，用户先获得确定性数据，再等待可选分析文本。图表与导出复用同一清洗结果，避免页面、CSV/XLSX 与自然语言结论出现权限或单位不一致。")

    add_heading(doc, "5 产品化、接口与前端", 1)
    add_heading(doc, "5.1 产品服务与会话模型", 2)
    add_paragraph(doc, "ProductQueryService 统一编排登录身份、会话、查询、结果、分享、导出、审计和告警。每个会话由 user_id+session_id 隔离；查询结果异步写回发起它的原始会话，避免用户快速切换线程后结果串线。SSE 在第一条状态消息中即返回 session_id，前端立即持久化，从而支持首轮澄清和断线后的正确续问。")
    add_paragraph(doc, "SQLite 保存会话元数据、待补全查询、审计事件、风险告警、冻结状态与分享记录；最近成功上下文以内存窗口加持久状态协作。前端提供会话列表、重命名/删除、历史详情、批量导出、分享和最终回答流。工作台采用约 29:71 的对话—报告布局，侧栏可在 240—480 px 间拖拽或折叠，报告区独立滚动，输入框固定，移动端自动转为纵向布局。")

    add_heading(doc, "5.2 OpenAPI 标准化", 2)
    add_paragraph(doc, "当前 OpenAPI 3.1 描述包含 29 条路径和 30 个操作，覆盖健康检查、认证、能力发现、同步/流式查询、最终回答、下钻、历史、导出、分享，以及管理端总览、指标、数据导入、审计、告警、冻结和用户管理。请求统一支持 request_id 与 caller_system，错误采用可枚举错误码，认证采用 HTTP Bearer。旧版 /query 与 /sessions 仍保留在描述中用于兼容提示，但实现按弃用策略返回 410。")
    api_domains = [
        ("认证与能力", "health、login/me/logout、capabilities"),
        ("问数与会话", "sync/stream query、final-answer stream、drill、session update/delete"),
        ("历史与交付", "history/detail/session history、export、share"),
        ("管理与治理", "overview、metrics、data import、audit、alerts、freezes、users"),
    ]
    add_table(doc, ["API 域", "代表能力"], api_domains, [1.7, 4.65])
    add_paragraph(doc, "接口可被数据中台、风控平台、营销平台和报表系统等 caller_system 复用。LLM 客户端与数据源通过抽象接口隔离，目前数据源适配器落地 DuckDB；这为后续数据库替换提供边界，但尚不代表已完成国产操作系统、国产 CPU、国产数据库或国密环境认证。")

    add_heading(doc, "5.3 数据增量发布", 2)
    add_paragraph(doc, "管理端数据更新采用“预览—校验—确认—事务发布”：上传 Excel 后先统计新增、未变化、将覆盖和文件内重复记录；若包含覆盖，必须二次确认；发布时获取数据库锁，并保证全量成功或回滚。该机制由早期“仅支持固定 Excel 全量重建”演进而来，使参赛系统具备持续维护数据的能力。")

    add_heading(doc, "6 安全治理与部署", 1)
    add_picture(doc, figs["security"], "图 4  纵深安全与数据治理模型", "同心层安全图，依次展示最小权限、可信执行、结果治理和审计处置，并注明系统边界。")
    add_heading(doc, "6.1 六类角色与最小权限", 2)
    roles = [
        ("总行管理", "21 项指标", "13 家机构", "经营全景"),
        ("分行管理", "21 项指标", "配置机构", "本机构经营"),
        ("业务条线", "ZB003—006、ZB018—021", "配置机构", "业务规模与效率"),
        ("风险岗位", "ZB013—017", "13 家机构", "资产质量与风险"),
        ("财务岗位", "ZB001—002、ZB007—012", "13 家机构", "财务、资本与盈利"),
        ("管理员", "无业务指标查询", "无业务结果", "用户、审计、告警、发布"),
    ]
    add_table(doc, ["角色", "指标范围", "机构范围", "职责"], roles, [1.2, 2.1, 1.45, 1.65])
    add_paragraph(doc, "授权发生在执行前，DuckDBExecutor 通过临时授权视图再次收缩数据域；SQLGuard 阻止通过自定义 SQL 绕过视图。当前事实表粒度是机构—指标—日期—数值，因此“行级权限”准确含义是机构记录级，而不是客户、手机号或身份证级权限。")

    add_heading(doc, "6.2 脱敏、导出、审计与告警", 2)
    add_paragraph(doc, "ZB013—017 等 S3 风险指标按角色动态脱敏；页面展示、自然语言答案、导出和分享使用同一清洗策略。单次导出超过 200 行需确认。审计事件以 previous_hash+event_hash 的 SHA-256 链记录，可识别中间篡改、链路断裂和尾部删除，但底层 SQLite 仍不是监管级不可变存储。")
    alerts = [
        ("S3 高频访问", "短时间连续查询敏感风险指标"),
        ("越权冻结", "10 分钟内 3 次越权，冻结 15 分钟"),
        ("单次大导出", "单次导出超过 200 行"),
        ("日累计导出", "单日累计超过 1,000 行"),
        ("非工作时段导出", "夜间/非工作时段导出 S3 数据"),
        ("分享频控", "10 分钟内创建 3 次分享"),
        ("宽查询", "单次查询超过 8 个指标"),
    ]
    add_table(doc, ["告警类型", "触发逻辑"], alerts, [1.9, 4.45])

    add_heading(doc, "6.3 容器化与公网验证", 2)
    add_paragraph(doc, "项目提供 Docker Compose 一键启动，后端、前端和 Caddy 三个服务协同运行，持久数据挂载至宿主机 deploy-data。容器内数据路径统一为 /app/source/dataset.xlsx、/app/data/bank_metrics.duckdb 和 /app/data/product.sqlite3，解决了 Windows 开发路径迁移到 Linux 容器时的配置问题。历史部署验证中，/health、首页、/docs 和 /openapi.json 均返回 200，端到端问数成功。")
    add_paragraph(doc, "公网版本曾在 https://shuheng.site 验证 HTTP 自动 308 跳转 HTTPS、TLS 1.3 和有效证书；Docker Compose 支持 Caddy 自动证书，也支持通过 compose.tls.yaml 使用自有证书。文档不记录测试账号密码等敏感部署信息。")
    add_callout(doc, "依赖风险", "历史前端 npm audit 报告包含 21 项依赖告警（1 低、4 中、16 高）。这不影响本轮功能验证，但在生产上线前必须升级或替换受影响依赖并重新构建、回归。", "FCE9E7", RED)

    add_heading(doc, "7 实验设计与量化结果", 1)
    add_heading(doc, "7.1 实验问题与指标", 2)
    add_paragraph(doc, "实验围绕四个问题展开：系统能否正确理解金融语义；多轮上下文能否稳定补全；规则与 LLM 混合路线能否兼顾准确率和延迟；工程与安全能力能否持续回归。主要指标如下。")
    metric_defs = [
        ("人工复核准确率", "人工判定为正确的问题数 / 总问题数；优先于简单字符串匹配"),
        ("运行错误率", "API、解析、SQL 或执行失败数 / 总问题数"),
        ("路由占比", "规则或 LLM 实际承担的问题比例"),
        ("延迟", "总耗时、平均值、中位数、P95、最大值"),
        ("上下文完成率", "多轮请求能否完成到可执行结果；不等同于语义准确率"),
        ("工程回归", "后端 pytest、前端构建/lint/测试、容器与 HTTP 可达性"),
    ]
    add_table(doc, ["指标", "定义"], metric_defs, [1.65, 4.7])

    add_heading(doc, "7.2 数据集与复核方法", 2)
    experiments = [
        ("单轮 200 题", "比赛/API 测试问题", "人工复核答案值、排序、单位与覆盖性"),
        ("多轮 73 题", "连续追问、指代、省略、跨轮补全", "按会话顺序执行并人工判定"),
        ("综合 273 题", "200 单轮 + 73 多轮", "统一记录 verdict、route、elapsed"),
        ("上下文 50 题", "历史缺陷回归集", "用于路由完成性、错误率与延迟；不作为最终准确率"),
        ("工程测试", "后端与前端测试套件", "在当前工作树执行并记录失败原因"),
    ]
    add_table(doc, ["实验集", "覆盖内容", "评价方式"], experiments, [1.5, 2.3, 2.55])

    add_heading(doc, "7.3 核心结果", 2)
    result_rows = [
        ("200 题首轮", "189 正确 / 10 不匹配 / 1 错误", "94.50%", "196 规则 / 3 LLM / 1 错误", "P95 0.08s"),
        ("200 题第二轮", "192 正确 / 8 不匹配 / 0 错误", "96.00%", "199 规则 / 1 LLM", "P95 0.09s"),
        ("73 题多轮", "70 正确 / 3 不匹配 / 0 错误", "95.89%", "70 规则 / 3 LLM", "综合报告子集"),
        ("273 题综合第二轮", "262 正确 / 11 不匹配 / 0 错误", "95.97%", "269 规则 / 4 LLM", "P95 0.11s"),
        ("273 题综合第三轮", "261 正确 / 12 不匹配 / 0 错误", "95.60%", "269 规则 / 4 LLM", "P95 0.10s"),
    ]
    add_table(doc, ["实验", "判定", "准确率", "路由", "延迟"], result_rows, [1.35, 2.15, 0.8, 1.35, 0.85])
    add_paragraph(doc, "第二轮综合报告是当前最完整的人工复核记录：273 题总耗时 105.74 秒，平均 0.387 秒，中位数 0.07 秒，P95 0.11 秒，最大值 20.03 秒。规则路径承担 98.53% 请求，说明低延迟由确定性主链提供；4 个 LLM 请求解释了最大时延与总耗时之间的长尾差异。第三轮准确率略降 0.37 个百分点，提示人工复核口径或个别修复可能产生波动，因此项目应保留逐题差异而非只看总分。")

    add_heading(doc, "7.4 多轮能力的演进性实验", 2)
    add_paragraph(doc, "早期 50 题多轮基线仅 13 题完全正确（26%），9 题核心正确但不完整（18%），22 题实质性错误（44%），6 题执行错误（12%）。错误集中在空槽位继承、日期回退、机构范围扩张、阈值方向、比例单位和排名方向。重构三态路由、PendingQuery 与答案覆盖校验后，同一类 50 题上下文回归实现 50/50 完成、0 SQL/解析/执行错误，总耗时约 44.27 秒，中位数 0.035 秒、P95 3.751 秒、最大 19.226 秒，其中约 4 次进入 LLM。")
    add_callout(doc, "不可直接比较", "“50/50 完成”只证明没有解析/执行失败，不能替代人工语义准确率。该历史结果中曾存在比例展示错误，后续已修复，因此最终成果采用 73 题多轮人工复核的 95.89% 作为准确率证据。", PALE_GOLD, GOLD)

    add_heading(doc, "7.5 当前工程验证", 2)
    qa_rows = [
        ("后端 pytest", "239 / 240 通过，14.14 秒", "唯一失败：测试固定期望 132,678 行；增量发布后实际 132,698 行"),
        ("前端构建", "通过", "vinext 构建成功，路由产物生成"),
        ("前端 lint", "通过", "当前代码风格检查无失败"),
        ("前端测试", "2 / 2 通过", "日期特性回归测试通过"),
        ("容器部署（历史）", "3 服务健康", "首页、健康检查、API 文档、OpenAPI 均返回 200"),
        ("端到端（历史）", "通过", "A 市存款查询成功返回 42.02"),
    ]
    add_table(doc, ["验证项", "结果", "说明"], qa_rows, [1.5, 1.55, 3.4])
    add_paragraph(doc, "当前后端通过率按用例数计算为 99.58%。失败项反映测试夹具与增量数据治理之间的耦合：测试应改为基于 build_metadata 或预期增量状态断言，而不应把事实表总行数硬编码为永久常量。该问题不应被隐瞒，也不应被误报为 NL2SQL 核心故障。")

    add_heading(doc, "7.6 误差分析", 2)
    error_types = [
        ("参考口径冲突", "6 个经营好坏问题与官方衍生规则不一致", "建立真值优先级与冲突清单，不为错误答案调参"),
        ("上下文污染", "错误继承上一轮指标、日期或机构", "三态路由 + PendingQuery + 成功/失败记忆隔离"),
        ("单位与比例", "0.0242% 被放大为 2.42% 或 242%", "编译与格式层显式单位换算，增加回归断言"),
        ("排名方向", "低值更优指标或 Bottom N 方向错误", "目录中存方向属性，候选排序状态每次重置"),
        ("省级口径", "“全行/全市/13个市”未归一", "扩展组织别名与全省聚合规则"),
        ("时间表达", "两位年份、自然日期无法解析", "日期解析器统一归一到精确 DATE"),
        ("LLM 尾延迟", "少数请求达 16—20 秒", "提高规则覆盖、设置超时与降级、记录分位数"),
    ]
    add_table(doc, ["误差类型", "表现", "治理策略"], error_types, [1.45, 2.35, 2.65])

    add_heading(doc, "8 技术攻关脉络", 1)
    add_picture(doc, figs["timeline"], "图 5  从原型到交付的技术演进时间线", "时间线从2026年7月21日到8月27日，展示原型、多轮重构、产品化、安全、评测和部署。")
    add_heading(doc, "8.1 阶段一：建立确定性骨架", 2)
    add_paragraph(doc, "首个阶段完成语义目录、规则规划、QueryPlan、SQL 编译、DuckDB 执行和答案生成，证明了“自然语言—结构化计划—确定性 SQL”的可行性。相较纯 LLM 生成，系统从一开始就保留了可解释计划和可测试 SQL。")
    add_heading(doc, "8.2 阶段二：用失败样本重构多轮", 2)
    add_paragraph(doc, "50 题多轮分析暴露了上下文继承的系统性缺陷。团队没有继续堆叠关键词规则，而是把“是否需要上下文”和“规则能否覆盖”拆为两个正交决策，并引入 PendingQuery、结果覆盖义务和持久状态。这一阶段完成了从“能对话”到“可控状态机”的转变。")
    add_heading(doc, "8.3 阶段三：补齐产品与安全闭环", 2)
    add_paragraph(doc, "产品化阶段引入登录、会话、流式查询、最终回答、导出、分享、管理后台和 OpenAPI；安全阶段进一步把授权前移到执行前，并实现机构/指标数据域、S3 脱敏、导出门槛、审计哈希链、风险告警和自动冻结。前端同步修复会话串线与首轮澄清 session_id 丢失，保证后端能力可被真实用户稳定使用。")
    add_heading(doc, "8.4 阶段四：评测驱动的口径收敛", 2)
    add_paragraph(doc, "200/273 题人工复核不只统计总分，还记录 route、elapsed、错误类型与标准答案冲突。修复集中在排名方向、全省口径、日期解析、比例单位和勾稽覆盖。团队明确采用官方衍生说明作为最高真值，拒绝为冲突答案做不可解释的特殊分支。")
    add_heading(doc, "8.5 阶段五：部署与持续数据更新", 2)
    add_paragraph(doc, "最终阶段形成 Docker Compose、Caddy HTTPS、持久卷、健康检查与公网验证，并把数据维护从全量重建扩展为管理端增量预览/发布。技术成果由“离线脚本”转化为可交互、可集成、可运维的系统。")

    add_heading(doc, "9 参赛成果、创新点与边界", 1)
    add_heading(doc, "9.1 可验证参赛成果", 2)
    outcomes = [
        ("算法效果", "273 题人工复核 95.97%，0 运行错误；多轮 73 题 95.89%"),
        ("效率", "规则主路径 98.53%；综合中位数 0.07 秒、P95 0.11 秒"),
        ("数据能力", "13 机构、21 指标、487 日期、132,698 事实；支持增量事务发布"),
        ("可信安全", "6 类角色、指标/机构域、动态脱敏、哈希链、7 类风险告警"),
        ("集成能力", "OpenAPI 3.1，29 路径、30 操作；同步/流式/管理接口"),
        ("交付能力", "Web 工作台、Docker Compose、Caddy HTTPS、持久化与健康检查"),
    ]
    add_table(doc, ["成果维度", "量化或可复现结果"], outcomes, [1.45, 4.9])
    add_callout(doc, "赛事信息边界", "当前仓库与同项目会话未检出官方名次、奖项或评委评分材料。若后续获得证书、排名或现场评分，应作为“官方赛事结果”独立补充，不能用内部准确率替代赛事名次。", PALE_GOLD, GOLD)

    add_heading(doc, "9.2 主要创新点", 2)
    innovations = [
        ("结构化可信中间层", "以 QueryPlan 连接自然语言、权限与 SQL，使语义可检查、SQL 可复现、错误可定位。"),
        ("双路规划、单一闸门", "规则与 LLM 分工而不分裂安全体系，所有路径共用授权、校验和结果义务。"),
        ("多轮状态正交化", "上下文依赖判断与规则覆盖判断解耦，PendingQuery 与成功记忆分离并持久化。"),
        ("答案覆盖验证", "从“SQL 能运行”提升到“用户问题的每个义务都被回答”。"),
        ("语义—安全一体化", "指标目录既用于理解，也用于角色白名单、方向判断、单位换算和脱敏。"),
        ("数据持续交付", "预览、冲突统计、覆盖确认、事务发布把静态比赛数据变为可治理数据服务。"),
    ]
    add_table(doc, ["创新点", "技术价值"], innovations, [1.75, 4.6])

    add_heading(doc, "9.3 当前边界与风险", 2)
    boundaries = [
        "准确率基于项目内部人工复核记录，尚未由赛事方或第三方独立复现；应保留逐题日志和复核人信息。",
        "当前数据源为单机 DuckDB，SQLite 承担状态与审计；不具备生产级高可用、灾备和监管级 WORM。",
        "机构级数据域不等同于客户级行权限；当前数据中没有客户身份字段，不应夸大为全量银行数据安全方案。",
        "容器化和标准接口提供适配基础，但未完成麒麟/UOS、鲲鹏/飞腾、国产数据库和国密 TLS 的兼容认证。",
        "LLM 路径存在 16—20 秒级尾延迟，应设置超时、缓存、降级策略并持续提升规则/检索覆盖。",
        "前端依赖审计存在历史高危项；正式上线前需要版本升级、SBOM、镜像扫描和供应链门禁。",
        "当前一个后端测试因固定行数断言失败；应修复测试夹具并在增量发布后自动触发完整回归。",
    ]
    for item in boundaries:
        add_bullet(doc, item)

    add_heading(doc, "9.4 下一阶段技术路线", 2)
    roadmap = [
        "P0：修复数据行数硬编码测试，固化 273 题逐题金标与冲突答案清单，形成可重复一键评测。",
        "P0：升级存在高危告警的前端依赖，补充镜像漏洞扫描与 secret 扫描。",
        "P1：建设指标口径管理、同义词、已验证问题样例检索和表字段召回，借鉴 SQLBot 的元数据治理思想，但不复制其代码或许可证受限实现。",
        "P1：为 LLMPlanner 增加超时、熔断、缓存、离线重放和模型版本对比，重点治理 P95/P99 尾延迟。",
        "P1：将审计链输出到对象锁定或专用审计存储，并增加备份、恢复演练和时间戳签名。",
        "P2：新增 PostgreSQL/国产数据库适配器与方言测试；在目标国产 OS/CPU/中间件环境完成正式兼容认证。",
        "P2：当业务引入人工审批、长任务恢复或多智能体协作时，再评估 LangGraph 等编排框架。",
    ]
    for item in roadmap:
        add_bullet(doc, item)

    add_heading(doc, "10 结论", 1)
    add_paragraph(doc, "BankInsight 通过“语义目录—QueryPlan—确定性编译—统一可信闸门—产品交付”的路线，将 NL2SQL 从模型演示提升为面向金融指标场景的可解释、可治理系统。273 题人工复核 95.97%、规则路径 98.53%、0 运行错误，以及当前 132,698 行数据、30 个 API 操作、6 类角色和完整容器部署，共同构成可验证的参赛技术成果。")
    add_paragraph(doc, "更重要的是，项目的技术攻关脉络清晰：多轮低准确率通过错误分析推动状态机重构；安全问题推动权限前移与审计闭环；静态数据推动增量发布；部署问题推动路径、持久化和 HTTPS 标准化。系统仍存在高可用、供应链、审计存储、尾延迟和国产化认证等边界，但这些边界已被明确量化，并转化为可执行的下一阶段路线。")

    add_heading(doc, "附录 A 核心代码模块索引", 1)
    modules = [
        ("text2sql/service.py", "NL2SQL 主编排、规则/LLM 尝试与可信执行"),
        ("text2sql/context_router.py", "三态上下文依赖判断与改写"),
        ("text2sql/rule_planner.py", "高频金融问题规则规划"),
        ("text2sql/llm_planner.py", "LLM 结构化规划与降级"),
        ("text2sql/sql_compiler.py", "QueryPlan 到确定性 DuckDB SQL"),
        ("text2sql/validators.py", "计划、SQL、语义、结果与答案覆盖校验"),
        ("text2sql/executor.py", "授权域内只读 DuckDB 执行"),
        ("text2sql/semantic_catalog.py", "指标、组织、单位、方向与衍生规则"),
        ("text2sql/product_service.py", "认证、查询、导出、分享、审计与风控编排"),
        ("text2sql/product_store.py", "会话、PendingQuery、审计哈希链、告警与冻结"),
        ("text2sql/metric_data_import.py", "Excel 增量预览、冲突识别与事务发布"),
        ("text2sql/api.py", "FastAPI/OpenAPI、SSE 和管理端接口"),
    ]
    add_table(doc, ["模块", "职责"], modules, [2.25, 4.1])

    add_heading(doc, "附录 B 关键量化口径与来源", 1)
    evidence = [
        ("数据规模", "当前 DuckDB 直接统计", "13 / 21 / 132,698 / 487"),
        ("API 规模", "当前 OpenAPI 文档统计", "29 paths / 30 operations"),
        ("综合准确率", "273 题第二轮人工复核 JSONL", "262 / 273 = 95.97%"),
        ("单轮准确率", "200 题第二轮人工复核 JSONL", "192 / 200 = 96.00%"),
        ("多轮准确率", "73 题综合报告子集", "70 / 73 = 95.89%"),
        ("路由占比", "273 题 route 字段", "269 / 273 = 98.53% 规则"),
        ("当前后端测试", "pytest -q", "239 / 240；1 个数据基线断言过时"),
        ("前端质量", "npm test/build/lint", "2 / 2；构建和 lint 通过"),
        ("部署验证", "历史会话的 Docker/Caddy 实测", "3 服务健康、HTTP/HTTPS/API 可达"),
    ]
    add_table(doc, ["指标", "来源类型", "结果"], evidence, [1.55, 2.75, 2.05])

    add_heading(doc, "附录 C 复现建议", 1)
    reproduction_steps = [
        "记录 Git 提交、配置快照、数据库 build_metadata、模型供应商与模型版本。",
        "从固定评测 JSONL 依次执行 200 题单轮和 73 题多轮，不打乱会话顺序。",
        "保存 question、session_id、route、QueryPlan、SQL 摘要、elapsed、result、verdict 和复核说明。",
        "自动汇总准确率、错误率、路由占比、平均/中位/P95/P99/最大延迟，并输出逐题差异。",
        "在增量数据发布、规则修改、模型升级和权限变更后自动运行后端、前端与 273 题回归。",
        "部署环境验证健康检查、登录、问数、流式回答、导出、分享、审计、告警和恢复流程。",
    ]
    for step in reproduction_steps:
        add_bullet(doc, step)

    doc.save(OUT_DOCX)


def build_markdown(figs: dict[str, Path]) -> None:
    rel = lambda p: p.relative_to(DOCS).as_posix()
    md = f"""# {REPORT_TITLE}

> {REPORT_SUBTITLE}  
> {VERSION}

## 执行摘要

BankInsight 是面向银行经营指标查询场景的垂直 NL2SQL 系统。它以金融语义目录和结构化 `QueryPlan` 为核心：高置信问题由规则规划器解析，复杂或未覆盖表达再由 LLM 补位；两条路径统一进入权限裁剪、SQL 安全、语义对齐和结果覆盖校验。

截至 2026-08-27，系统覆盖 **13 家机构、21 项指标、487 个日期、132,698 条指标事实**；提供 **29 条 API 路径、30 个 OpenAPI 操作和 6 类业务角色**。人工复核的 273 题综合集达到 **262 题正确、11 题不匹配、0 运行错误，准确率 95.97%**；规则主路径覆盖 269 题，占 **98.53%**。

![评测结果与工程质量快照]({rel(figs['metrics'])})

> **证据口径**：当前状态以 2026-08-27 工作树为准。项目与同项目会话中未发现官方获奖名次或证书，因此本文只报告可验证技术成果，不虚构赛事排名。

## 1. 项目概述与问题定义

### 1.1 业务目标

比赛数据由银行机构、指标、日期和数值构成。用户问题表面上是自然语言转 SQL，实际同时包含金融术语识别、时间归一、组织范围、统计动作、排名方向、单位换算、衍生指标和权限边界。系统目标是让业务人员无需编写 SQL，即可完成点查、对比、排名、趋势、结构、环比/同比和勾稽分析，并获得表格、图表与自然语言结论。

### 1.2 核心挑战

| 挑战 | 技术含义 |
|---|---|
| 语义歧义 | “最高/最低”“全省/全行”“今年2月1日”等表达必须映射为可审计槽位。 |
| 多轮依赖 | 既要正确继承省略信息，又不能让上一轮状态污染当前问题。 |
| 精确计算 | 排名方向、比例单位、同比/环比、日均、勾稽需确定性执行。 |
| 金融安全 | 不同岗位只能查询授权机构与指标，结果、导出和分享必须一致脱敏。 |
| 工程交付 | 能力需通过 API、Web、容器、HTTPS、数据发布和运维审计真正可用。 |

### 1.3 设计原则

- 可信优先：任何 SQL 都必须经过统一校验与授权执行。
- 确定性优先：事实值和计算由规则、编译器与数据库完成，模型不口算。
- 渐进增强：以有限状态机解决当前问题，复杂度确有需要时再引入重型编排。
- 证据驱动：用人工复核、错误率、路由、延迟分位数、测试和部署量化成果。
- 边界清晰：容器化不等于国产化认证，哈希链不等于监管级 WORM。

## 2. 数据与金融语义底座

### 2.1 当前数据规模

| 维度 | 当前值 | 说明 |
|---|---:|---|
| 机构数 | 13 | 组织编码与名称双重识别 |
| 指标数 | 21 | 资产、负债、收入、风险、资本等主题 |
| 衍生规则 | 10 | 组合、排名、阈值、增速、勾稽等 |
| 指标事实 | 132,698 行 | 较历史基线增加 20 条增量记录 |
| 日期范围 | 2024-12-31—2026-06-30 | 487 个不同日期 |
| 主键 | 日期+机构+指标 | 当前重复主键为 0 |

数据层使用 DuckDB；增量更新采用“预览—冲突统计—覆盖确认—事务发布”。增量文件可以只含“指标数据表”并包含较少的已有指标，但不能通过文件新增指标定义。若使用未同步的原始 Excel 全量重建，可能覆盖增量数据，生产流程必须同步归档和源文件。

### 2.2 语义目录与口径优先级

`SemanticCatalog` 管理指标编码、别名、单位、方向属性、组织别名和派生规则。对“经营情况好/差”，系统按官方衍生说明：多数指标越高越好，ZB012、ZB013、ZB017 越低越好；13 家机构中 `RANK` 1—3 为好、4—9 中性、10—13 为差。

测试材料中第 91、93、151、152、191、192 题的参考答案与官方衍生规则冲突。当前实现坚持官方规则，并把此类问题标记为“口径冲突”，不为错误答案定向调参。

### 2.3 QueryPlan 可信中间表示

`QueryPlan` 把自然语言拆为指标集合、机构集合、时间范围、查询形态、排序方向、Top/Bottom N、阈值、聚合方式、单位和派生计算。支持点查、多指标、排名、省均、阈值、期间值、日均、求和、结构占比、同比/环比、勾稽、趋势和画像。

## 3. 总体技术路线

![BankInsight 总体技术架构]({rel(figs['architecture'])})

### 3.1 规则优先、模型补位

高频金融问数具有强结构性。系统优先用 `RulePlanner` 生成高置信 `QueryPlan`，再由 `SQLCompiler` 确定性生成 DuckDB SQL；规则未覆盖、表达歧义或复杂改写才进入 `LLMPlanner`。273 题中规则路径覆盖 269 题（98.53%），使整体中位延迟约 0.07 秒；少数 LLM 请求形成 20 秒级尾延迟。

### 3.2 全链路可信闸门

![单次问数可信执行链]({rel(figs['pipeline'])})

执行顺序为：身份与请求 → 上下文决策 → 槽位补全 → 规则/模型规划 → 确定性编译 → 权限裁剪 → 计划/SQL/语义校验 → 只读执行 → 结果/答案覆盖校验 → 图表与解释 → 审计闭环。任何来源的 SQL 都不能绕过同一闸门。

## 4. 模型架构与关键技术

### 4.1 三态上下文路由

`ContextRouter` 先判断问题为 `SELF_CONTAINED`、`CONTEXT_REQUIRED` 或 `UNCERTAIN`，再决定是否读取历史。当前显式槽位优先级最高，其次是 `PendingQuery`、最近成功上下文，最后才是 LLM 推断。该设计解决早期空槽位继承导致的日期、机构和指标污染。

### 4.2 PendingQuery 与持久化

缺机构、日期或指标时，失败请求不写入成功记忆，而是生成待补全状态。`PendingQuery` 按 `user_id+session_id` 写入 SQLite，后端重启后仍能续接。历史复测中，服务重启后补充“J市”，仍正确返回净利润 175.6 与 +25.34。

### 4.3 规则规划、LLM 规划与上下文最小化

规则覆盖点查、排名、阈值、期间、结构、同比/环比、勾稽、趋势和画像；日期解析支持“今年2月1日”“25年2月1日”，组织解析覆盖“全省/全行/全市/13个市”。LLM 只接收当前问题、必要成功历史、结构化槽位和必要结果摘要，不接收完整 SQL、完整日志或失败执行内容；输出必须转成 `QueryPlan` 并通过统一校验。

### 4.4 确定性编译与四级验证

- `PlanValidator`：检查槽位、查询形态、阈值、排名、时间和组织约束。
- `SQLGuard`：基于 sqlglot AST 仅允许只读 SQL，限制表、字段、函数和危险语句。
- `AlignmentValidator`：确保 SQL 与计划的指标、机构、日期、聚合和排序一致。
- `ResultValidator`：检查结果结构、空值、类型、排序和单位。
- `AnswerCoverageValidator`：保证排名、差值、比较、阈值、计数、勾稽等用户义务被完整回答。

确定性编译显式处理低值更优指标、Bottom N、比例单位、基期定位和窗口排名，避免模型把 0.0242% 放大为 2.42% 或 242%。

### 4.5 为什么当前未使用 LangGraph

当前流程是有限、可枚举的状态机，尚不需要多智能体、人工审批或复杂中断恢复。引入 LangGraph 会增加状态、调试和部署复杂度。后续出现人工审批、长任务恢复或多智能体协作时再评估。

## 5. 产品化、接口与前端

### 5.1 产品服务与会话

`ProductQueryService` 统一编排登录、会话、查询、结果、分享、导出、审计和告警。会话以 `user_id+session_id` 隔离；异步结果写回原始线程，防止用户切换线程后串线。SSE 第一条状态消息即携带 `session_id`，前端立即持久化，支持首轮澄清和断线续问。

Web 工作台包含会话列表、历史详情、流式回答、表格/图表、导出、分享和管理端。桌面端约 29:71 对话—报告布局，侧栏可在 240—480px 拖拽或折叠；报告独立滚动，输入框固定；移动端自适应为纵向布局。

### 5.2 OpenAPI 标准化

当前 OpenAPI 3.1 包含 **29 paths / 30 operations**，覆盖健康、认证、能力发现、同步/流式查询、最终回答、下钻、历史、导出、分享，以及管理端总览、指标、数据导入、审计、告警、冻结与用户管理。请求统一支持 `request_id` 与 `caller_system`，认证使用 HTTP Bearer，错误采用枚举错误码。旧版 `/query` 与 `/sessions` 仅作弃用兼容并返回 410。

接口面向数据中台、风控平台、营销平台、报表系统等复用；当前数据源适配器实现 DuckDB。抽象层为数据库替换提供接口，但不代表已完成国产 OS/CPU/数据库或国密认证。

## 6. 安全治理与部署

![纵深安全与数据治理]({rel(figs['security'])})

### 6.1 六类角色

| 角色 | 指标范围 | 机构范围 | 职责 |
|---|---|---|---|
| 总行管理 | 21 项 | 13 家 | 全景经营 |
| 分行管理 | 21 项 | 配置机构 | 本机构经营 |
| 业务条线 | ZB003—006、ZB018—021 | 配置机构 | 业务规模与效率 |
| 风险岗位 | ZB013—017 | 13 家 | 资产质量与风险 |
| 财务岗位 | ZB001—002、ZB007—012 | 13 家 | 财务、资本与盈利 |
| 管理员 | 无业务查询 | 无业务结果 | 用户、审计、告警、发布 |

授权在执行前完成，DuckDB 临时授权视图再次收缩数据域；`SQLGuard` 阻止绕过。当前“行级权限”准确含义是机构记录级，不是客户级。

### 6.2 脱敏、审计和七类告警

ZB013—017 等 S3 指标按角色动态脱敏，页面、回答、导出和分享使用同一清洗策略。审计事件通过 `previous_hash + event_hash` 的 SHA-256 链检测篡改、链断和尾删，但 SQLite 不是监管级 WORM。

七类告警包括：S3 高频访问；10 分钟 3 次越权后冻结 15 分钟；单次导出超过 200 行；单日超过 1,000 行；非工作时段导出 S3；10 分钟 3 次分享；单查询超过 8 个指标。

### 6.3 容器与 HTTPS

Docker Compose 一键启动后端、前端和 Caddy；数据挂载到宿主机 `deploy-data`。容器内路径统一为 `/app/source/dataset.xlsx`、`/app/data/bank_metrics.duckdb`、`/app/data/product.sqlite3`。历史公网验证中，首页、`/health`、`/docs`、`/openapi.json` 均返回 200，HTTP 308 跳转 HTTPS，TLS 1.3 可用。

> **供应链风险**：历史 `npm audit` 有 21 项告警（1 低、4 中、16 高），正式上线前需要依赖升级、镜像扫描和完整回归。

## 7. 实验设计与量化结果

### 7.1 实验设计

| 实验集 | 规模 | 目标 | 评价方式 |
|---|---:|---|---|
| 单轮 | 200 | 金融语义、计算、排序、单位 | 人工复核答案值与覆盖 |
| 多轮 | 73 | 指代、省略、补全、跨轮状态 | 按会话顺序执行并人工复核 |
| 综合 | 273 | 统一效果、路由和时延 | `verdict + route + elapsed` |
| 上下文回归 | 50 | 历史缺陷完成性与错误率 | 不作为最终语义准确率 |
| 工程测试 | 240 后端 + 前端 | 回归与交付质量 | pytest、build、lint、test |

主要指标为人工复核准确率、运行错误率、规则/LLM 路由占比、总耗时、平均/中位/P95/最大延迟、上下文完成率和工程测试结果。

### 7.2 核心结果

| 实验 | 正确/不匹配/错误 | 准确率 | 路由 | 延迟 |
|---|---|---:|---|---|
| 200 题首轮 | 189 / 10 / 1 | 94.50% | 196 规则 / 3 LLM / 1 错误 | P95 0.08s |
| 200 题第二轮 | 192 / 8 / 0 | 96.00% | 199 规则 / 1 LLM | P95 0.09s |
| 73 题多轮 | 70 / 3 / 0 | 95.89% | 70 规则 / 3 LLM | 综合报告子集 |
| 273 题综合第二轮 | 262 / 11 / 0 | **95.97%** | 269 规则 / 4 LLM | **P95 0.11s** |
| 273 题综合第三轮 | 261 / 12 / 0 | 95.60% | 269 规则 / 4 LLM | P95 0.10s |

第二轮综合报告总耗时 105.74 秒，平均 0.387 秒，中位数 0.07 秒，P95 0.11 秒，最大 20.03 秒。规则路径承担 98.53% 请求；4 个 LLM 请求解释了尾延迟。

### 7.3 多轮演进

早期 50 题仅 13 题完全正确（26%）、9 题核心正确（18%）、22 题实质错误（44%）、6 题运行错误（12%）。重构后同类回归集 50/50 完成、0 SQL/解析/执行错误，总耗时 44.27 秒、中位 0.035 秒、P95 3.751 秒、最大 19.226 秒，约 4 次使用 LLM。

“50/50 完成”不等于语义准确率；该历史集曾存在比例展示错误，后续已修复。最终多轮准确率以 73 题人工复核的 95.89% 为准。

### 7.4 当前工程验证

| 验证项 | 结果 | 说明 |
|---|---|---|
| 后端 pytest | 239 / 240 通过，14.14s | 失败项固定期望 132,678 行；当前实际 132,698 行 |
| 前端构建 | 通过 | vinext 构建成功 |
| 前端 lint | 通过 | 无失败 |
| 前端测试 | 2 / 2 通过 | 日期特性回归通过 |
| 容器部署（历史） | 3 服务健康 | 首页、健康、API 文档、OpenAPI 可达 |
| 端到端（历史） | 通过 | A 市存款查询返回 42.02 |

后端按用例数通过率为 99.58%。唯一失败是测试夹具未适应 20 条增量数据，不是 NL2SQL 查询逻辑失败，但应尽快移除总行数硬编码并在数据发布后自动回归。

## 8. 技术攻关脉络

![技术演进时间线]({rel(figs['timeline'])})

1. **确定性骨架**：完成语义目录、QueryPlan、规则规划、SQL 编译和 DuckDB 执行。
2. **多轮重构**：用 50 题失败分类推动三态路由、PendingQuery、成功/失败记忆隔离。
3. **产品与安全**：补齐登录、会话、流式回答、导出、分享、OpenAPI、角色数据域、脱敏、审计、告警和冻结。
4. **评测收敛**：通过 200/273 题人工复核修复排名方向、省级口径、日期、比例和勾稽覆盖，并建立真值优先级。
5. **部署与更新**：形成 Docker+Caddy+HTTPS+持久卷，并把数据维护从全量重建扩展为增量事务发布。

## 9. 参赛成果、创新点与边界

### 9.1 可验证成果

- 算法：273 题人工复核 95.97%，0 运行错误；73 题多轮 95.89%。
- 效率：规则路径 98.53%；综合中位 0.07 秒，P95 0.11 秒。
- 数据：13 机构、21 指标、487 日期、132,698 行，支持增量事务发布。
- 安全：6 类角色、指标/机构域、动态脱敏、哈希链、7 类告警。
- 集成：OpenAPI 3.1，29 路径、30 操作；支持同步、流式和管理接口。
- 交付：Web 工作台、Docker Compose、Caddy HTTPS、持久化和健康检查。

### 9.2 创新点

- **结构化可信中间层**：QueryPlan 连接语言、权限与 SQL。
- **双路规划、单一闸门**：规则与 LLM 分工但共用安全和校验。
- **多轮状态正交化**：上下文依赖与规则覆盖解耦，PendingQuery 持久化。
- **答案覆盖校验**：从 SQL 可执行提升到用户义务完整回答。
- **语义—安全一体化**：同一指标目录驱动别名、方向、单位、白名单和脱敏。
- **数据持续交付**：预览、冲突统计、覆盖确认和事务发布形成数据闭环。

### 9.3 边界与风险

- 内部人工复核尚未由赛事方或第三方独立复现。
- DuckDB 与 SQLite 是单机架构，不具备生产级高可用、灾备和 WORM。
- 当前是机构记录级权限，不是客户级权限。
- 未完成国产 OS/CPU/数据库与国密环境认证。
- LLM 路径存在 16—20 秒级尾延迟。
- 前端依赖存在历史高危告警。
- 当前一个后端测试因固定行数断言失败。

### 9.4 下一阶段

1. 修复数据行数硬编码测试，固化 273 题逐题金标和冲突清单。
2. 升级前端依赖，增加 SBOM、secret 和镜像漏洞扫描。
3. 建设指标口径、同义词、已验证样例检索和表字段召回；借鉴 SQLBot 元数据治理思想，但不复制受限代码。
4. 为 LLM 增加超时、熔断、缓存、离线重放与模型版本对比。
5. 把审计链写入对象锁定或专用审计存储，补充备份恢复演练。
6. 增加 PostgreSQL/国产数据库适配器和目标环境兼容认证。
7. 只有业务出现人工审批、长任务恢复或多智能体协作时再评估 LangGraph。

## 10. 结论

BankInsight 通过“语义目录—QueryPlan—确定性编译—统一可信闸门—产品交付”，把 NL2SQL 从模型演示提升为面向金融指标的可解释、可治理系统。273 题人工复核 95.97%、规则路径 98.53%、0 运行错误，以及 132,698 行数据、30 个 API 操作、6 类角色和容器化部署，共同构成可验证的参赛技术成果。

## 附录 A：核心模块索引

| 模块 | 职责 |
|---|---|
| `text2sql/service.py` | NL2SQL 主编排、规则/LLM 尝试与可信执行 |
| `text2sql/context_router.py` | 三态上下文依赖判断 |
| `text2sql/rule_planner.py` | 高频金融问题规则规划 |
| `text2sql/llm_planner.py` | LLM 结构化规划与降级 |
| `text2sql/sql_compiler.py` | QueryPlan 到 DuckDB SQL |
| `text2sql/validators.py` | 计划、SQL、语义、结果和答案覆盖校验 |
| `text2sql/executor.py` | 授权域内只读执行 |
| `text2sql/semantic_catalog.py` | 指标、组织、单位、方向和衍生规则 |
| `text2sql/product_service.py` | 产品、安全与审计编排 |
| `text2sql/product_store.py` | 会话、PendingQuery、审计、告警与冻结 |
| `text2sql/metric_data_import.py` | 增量预览与事务发布 |
| `text2sql/api.py` | FastAPI/OpenAPI 与 SSE 接口 |

## 附录 B：复现建议

1. 固定 Git 提交、配置、数据库 `build_metadata`、模型供应商和模型版本。
2. 按顺序执行 200 题单轮与 73 题多轮，保留会话边界。
3. 保存问题、会话、路由、QueryPlan、SQL 摘要、耗时、结果、判定和说明。
4. 自动汇总准确率、错误率、路由占比、平均/中位/P95/P99/最大延迟及逐题差异。
5. 在数据发布、规则修改、模型升级和权限变更后运行后端、前端与 273 题回归。
6. 在部署环境验证健康、登录、问数、流式回答、导出、分享、审计、告警和恢复流程。
"""
    OUT_MD.write_text(md, encoding="utf-8")


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    figs = {
        "architecture": create_architecture_figure(),
        "pipeline": create_pipeline_figure(),
        "timeline": create_timeline_figure(),
        "metrics": create_metrics_figure(),
        "security": create_security_figure(),
    }
    build_markdown(figs)
    build_docx(figs)
    print(f"Markdown: {OUT_MD}")
    print(f"DOCX: {OUT_DOCX}")
    for key, value in figs.items():
        print(f"Figure {key}: {value}")


if __name__ == "__main__":
    main()
