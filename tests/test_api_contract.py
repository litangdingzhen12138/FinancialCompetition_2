from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from text2sql import api
from text2sql.datasource import DataSourceAdapter, DuckDBDataSourceAdapter
from text2sql.product_models import (
    ChartSpec,
    DataColumn,
    Insight,
    ProductQueryResponse,
)


def _query_response(session_id: str) -> ProductQueryResponse:
    return ProductQueryResponse(
        api_version="v1",
        query_id="query-001",
        session_id=session_id,
        status="completed",
        question="查询存款余额",
        route="rule",
        answer_mode="rule",
        answer_status="completed",
        generated_at="2026-08-23T00:00:00+00:00",
        duration_ms=12,
        plan={"operation": "value"},
        sql="SELECT 1",
        columns=(DataColumn("metric_value", "指标值", "number", "亿元"),),
        records=({"metric_value": 42.32},),
        row_count=1,
        truncated=False,
        visualization=ChartSpec("metric", "存款余额"),
        alternatives=(),
        insight=Insight("查询完成", "存款余额为42.32亿元"),
    )


def test_query_contract_returns_request_metadata(monkeypatch) -> None:
    class StubProductService:
        def query(self, question, session_id, user):
            assert question == "查询存款余额"
            assert user.role == "analyst"
            return _query_response(session_id)

    monkeypatch.setattr(api, "get_product_service", lambda: StubProductService())
    client = TestClient(api.app)

    response = client.post(
        "/api/v1/queries",
        headers={"X-Request-ID": "header-request"},
        json={
            "question": "查询存款余额",
            "session_id": "integration-session",
            "request_id": "body-request-001",
            "caller_system": "risk-platform",
        },
    )

    assert response.status_code == 200
    assert response.headers["X-Request-ID"] == "body-request-001"
    assert response.json()["request_id"] == "body-request-001"
    assert response.json()["caller_system"] == "risk-platform"


def test_validation_error_keeps_detail_and_adds_machine_fields() -> None:
    client = TestClient(api.app)

    response = client.post(
        "/api/v1/queries",
        headers={"X-Request-ID": "validation-request-001"},
        json={"question": ""},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "REQUEST_VALIDATION_ERROR"
    assert response.json()["detail"]
    assert response.json()["request_id"] == "validation-request-001"
    assert response.headers["X-Request-ID"] == "validation-request-001"


def test_query_error_uses_body_request_id(monkeypatch) -> None:
    class FailingProductService:
        def query(self, question, session_id, user):
            raise ValueError("查询条件不完整")

    monkeypatch.setattr(api, "get_product_service", lambda: FailingProductService())
    client = TestClient(api.app)

    response = client.post(
        "/api/v1/queries",
        json={
            "question": "查询数据",
            "request_id": "body-error-request-001",
        },
    )

    assert response.status_code == 400
    assert response.json() == {
        "code": "REQUEST_INVALID",
        "detail": "查询条件不完整",
        "request_id": "body-error-request-001",
    }
    assert response.headers["X-Request-ID"] == "body-error-request-001"


def test_openapi_exposes_query_contract_and_bearer_scheme() -> None:
    api.app.openapi_schema = None
    schema = api.app.openapi()

    assert schema["components"]["securitySchemes"]["BearerAuth"] == {
        "type": "http",
        "scheme": "bearer",
        "bearerFormat": "opaque-token",
    }
    operation = schema["paths"]["/api/v1/queries"]["post"]
    assert operation["security"] == [{"BearerAuth": []}]
    assert operation["responses"]["200"]["content"]["application/json"][
        "schema"
    ] == {"$ref": "#/components/schemas/ProductQueryResponse"}
    request_fields = schema["components"]["schemas"]["QueryRequest"]["properties"]
    assert {"question", "session_id", "request_id", "caller_system"} <= set(
        request_fields
    )


def test_default_service_uses_declared_data_source_boundary(service) -> None:
    assert isinstance(service.data_source, DuckDBDataSourceAdapter)
    assert isinstance(service.data_source, DataSourceAdapter)


def test_query_xlsx_export_uses_chinese_headers(monkeypatch) -> None:
    class StubProductService:
        def export_record(self, query_id, user, *, confirm_large_export):
            assert query_id == "query-ratio"
            assert user.can_export
            assert confirm_large_export is False
            return {
                "question": "查询存贷比",
                "answer": "存贷比为82.21%",
                "created_at": "2026-08-27T10:00:00+08:00",
                "columns": ["org_name", "numerator", "denominator", "derived_value"],
                "rows": [["江苏省J市农商行", 47.5, 57.78, 82.2084]],
                "plan": {},
            }

        def export_headers(self, record):
            assert record["columns"][1] == "numerator"
            return ("机构名称", "各项贷款余额", "各项存款余额", "存贷比")

    monkeypatch.setattr(api, "get_product_service", lambda: StubProductService())
    response = TestClient(api.app).get("/api/v1/queries/query-ratio/export")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert tuple(workbook["查询结果"].iter_rows(min_row=5, max_row=5, values_only=True))[0] == (
        "机构名称",
        "各项贷款余额",
        "各项存款余额",
        "存贷比",
    )


def test_admin_metric_import_contract_uses_raw_xlsx_body(monkeypatch) -> None:
    class StubProductService:
        def admin_metrics(self, user):
            assert user.can_view_admin
            return [
                {
                    "metric_id": "ZB001",
                    "metric_name": "各项存款余额",
                    "description": "各项存款余额",
                    "unit": "亿元",
                }
            ]

        def preview_metric_data_import(self, content, filename, user):
            assert user.can_view_admin
            assert content == b"xlsx-content"
            assert filename == "increment.xlsx"
            return {"valid": True, "insert_count": 1, "overwrite_count": 0}

        def publish_metric_data_import(
            self, content, filename, user, *, confirm_overwrite
        ):
            assert user.can_view_admin
            assert content == b"xlsx-content"
            assert filename == "increment.xlsx"
            assert confirm_overwrite is True
            return {"published": True, "insert_count": 1, "overwrite_count": 0}

    monkeypatch.setattr(api, "get_product_service", lambda: StubProductService())
    client = TestClient(api.app)

    metrics = client.get("/api/v1/admin/metrics")
    preview = client.post(
        "/api/v1/admin/data-imports/preview?filename=increment.xlsx",
        content=b"xlsx-content",
    )
    publish = client.post(
        "/api/v1/admin/data-imports/publish?filename=increment.xlsx&confirm_overwrite=true",
        content=b"xlsx-content",
    )

    assert metrics.status_code == 200
    assert metrics.json()["total"] == 1
    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert publish.status_code == 200
    assert publish.json()["published"] is True
