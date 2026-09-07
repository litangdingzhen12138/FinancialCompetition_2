from __future__ import annotations

from io import BytesIO
import shutil

import duckdb
from openpyxl import Workbook
import pytest

from text2sql.metric_data_import import (
    MetricDataImporter,
    OverwriteConfirmationRequired,
)
from text2sql.semantic_catalog import SemanticCatalog


def _workbook_bytes(
    rows: list[tuple[object, ...]],
    *,
    metric_rows: list[tuple[object, ...]] | None = None,
    data_sheet_title: str = "指标数据表",
) -> bytes:
    workbook = Workbook()
    if metric_rows is None:
        data_sheet = workbook.active
        data_sheet.title = data_sheet_title
    else:
        metric_sheet = workbook.active
        metric_sheet.title = "指标清单表"
        metric_sheet.append(("指标编号", "指标名称", "指标含义", "指标单位"))
        for row in metric_rows:
            metric_sheet.append(row)
        data_sheet = workbook.create_sheet("指标数据表")
    data_sheet.append(("数据日期", "指标编号", "指标名称", "机构编号", "指标值"))
    for row in rows:
        data_sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.fixture
def importer(service, tmp_path) -> MetricDataImporter:
    db_path = tmp_path / "metric-import.duckdb"
    shutil.copy2(service.db_path, db_path)
    return MetricDataImporter(db_path)


def test_data_only_workbook_inserts_subset_and_survives_catalog_reload(importer) -> None:
    content = _workbook_bytes(
        [("2030-01-01", "ZB001", "各项存款余额", "ORG001", 123.45)]
    )

    preview = importer.preview(content, "increment.xlsx")
    published = importer.publish(content, "increment.xlsx", confirm_overwrite=False)

    assert preview["valid"] is True
    assert preview["upload_mode"] == "metric_data_only"
    assert preview["insert_count"] == 1
    assert published["published"] is True
    assert SemanticCatalog(importer.db_path).metrics["ZB001"].name == "各项存款余额"
    with duckdb.connect(str(importer.db_path), read_only=True) as connection:
        value = connection.execute(
            "SELECT metric_value FROM metric_values "
            "WHERE data_date = '2030-01-01' AND metric_id = 'ZB001' AND org_id = 'ORG001'"
        ).fetchone()[0]
    assert value == pytest.approx(123.45)


def test_single_sheet_workbook_does_not_require_data_sheet_name(importer) -> None:
    content = _workbook_bytes(
        [("2030-01-01", "ZB001", "各项存款余额", "ORG001", 123.45)],
        data_sheet_title="Sheet1",
    )

    preview = importer.preview(content, "single-sheet.xlsx")

    assert preview["valid"] is True
    assert preview["upload_mode"] == "metric_data_only"
    assert preview["insert_count"] == 1


def test_multiple_sheet_workbook_still_requires_data_sheet_name(importer) -> None:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Sheet1"
    data_sheet.append(("数据日期", "指标编号", "指标名称", "机构编号", "指标值"))
    data_sheet.append(("2030-01-01", "ZB001", "各项存款余额", "ORG001", 123.45))
    workbook.create_sheet("说明")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ValueError, match="缺少.*指标数据表.*工作表"):
        importer.preview(output.getvalue(), "multiple-sheets.xlsx")


def test_complete_workbook_allows_metric_catalog_subset(importer) -> None:
    content = _workbook_bytes(
        [("2030-01-02", "ZB013", "不良贷款率", "ORG001", 1.25)],
        metric_rows=[("ZB013", "不良贷款率", "上传说明可省略", "%")],
    )

    preview = importer.preview(content, "complete.xlsx")

    assert preview["valid"] is True
    assert preview["upload_mode"] == "complete_workbook"
    assert preview["metric_ids"] == ["ZB013"]


def test_unknown_metric_and_name_mismatch_are_rejected(importer) -> None:
    content = _workbook_bytes(
        [
            ("2030-01-01", "ZB022", "新增指标", "ORG001", 1),
            ("2030-01-01", "ZB001", "错误名称", "ORG001", 2),
        ]
    )

    preview = importer.preview(content, "unknown.xlsx")

    assert preview["valid"] is False
    assert preview["can_publish"] is False
    assert any("当前指标清单之外" in item for item in preview["errors"])
    assert any("指标名称与当前清单不一致" in item for item in preview["errors"])


def test_file_duplicates_are_deduplicated_or_rejected(importer) -> None:
    same = _workbook_bytes(
        [
            ("2030-01-03", "ZB001", "各项存款余额", "ORG001", 10),
            ("2030-01-03", "ZB001", "各项存款余额", "ORG001", 10),
        ]
    )
    conflicting = _workbook_bytes(
        [
            ("2030-01-04", "ZB001", "各项存款余额", "ORG001", 10),
            ("2030-01-04", "ZB001", "各项存款余额", "ORG001", 11),
        ]
    )

    same_preview = importer.preview(same, "same.xlsx")
    conflict_preview = importer.preview(conflicting, "conflicting.xlsx")

    assert same_preview["valid"] is True
    assert same_preview["duplicate_count"] == 1
    assert same_preview["parsed_row_count"] == 1
    assert conflict_preview["valid"] is False
    assert conflict_preview["conflicting_duplicate_count"] == 1


def test_existing_value_requires_confirmation_before_overwrite(importer) -> None:
    with duckdb.connect(str(importer.db_path), read_only=True) as connection:
        data_date, metric_id, org_id, current = connection.execute(
            "SELECT data_date, metric_id, org_id, metric_value FROM metric_values LIMIT 1"
        ).fetchone()
        metric_name = connection.execute(
            "SELECT metric_name FROM metrics WHERE metric_id = ?", [metric_id]
        ).fetchone()[0]
    replacement = float(current) + 1.0
    content = _workbook_bytes(
        [(str(data_date), str(metric_id), str(metric_name), str(org_id), replacement)]
    )

    preview = importer.preview(content, "overwrite.xlsx")
    with pytest.raises(OverwriteConfirmationRequired):
        importer.publish(content, "overwrite.xlsx", confirm_overwrite=False)
    result = importer.publish(content, "overwrite.xlsx", confirm_overwrite=True)

    assert preview["overwrite_count"] == 1
    assert result["overwrite_count"] == 1
    assert result["overwrite_details"][0]["old_value"] == pytest.approx(float(current))
    assert result["overwrite_details"][0]["new_value"] == pytest.approx(replacement)
