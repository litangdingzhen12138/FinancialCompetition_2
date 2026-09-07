"""Validate and merge administrator-uploaded metric workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import hashlib
from io import BytesIO
import math
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import load_workbook

from .database_lock import database_lock


DATA_SHEET = "指标数据表"
METRIC_SHEET = "指标清单表"
DATA_HEADERS = ("数据日期", "指标编号", "指标名称", "机构编号", "指标值")
METRIC_HEADERS = ("指标编号", "指标名称", "指标含义", "指标单位")
MAX_VALIDATION_ERRORS = 50


@dataclass(frozen=True, slots=True)
class ImportRow:
    data_date: str
    metric_id: str
    metric_name: str
    org_id: str
    metric_value: float

    @property
    def key(self) -> tuple[str, str, str]:
        return (self.data_date, self.metric_id, self.org_id)


class OverwriteConfirmationRequired(ValueError):
    def __init__(self, preview: dict[str, Any]) -> None:
        super().__init__(f"本次发布将覆盖{preview['overwrite_count']}条现有数据，需要确认")
        self.preview = preview


def _same_value(left: float, right: float) -> bool:
    return math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-12)


def _date_value(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    candidates = (
        text,
        text.replace("/", "-"),
        text.replace("年", "-").replace("月", "-").replace("日", ""),
    )
    for candidate in candidates:
        try:
            return datetime.fromisoformat(candidate).date().isoformat()
        except ValueError:
            try:
                return date.fromisoformat(candidate).isoformat()
            except ValueError:
                continue
    return None


def _number_value(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _append_error(errors: list[str], message: str) -> None:
    if len(errors) < MAX_VALIDATION_ERRORS:
        errors.append(message)


class MetricDataImporter:
    def __init__(self, db_path: Path) -> None:
        self.db_path = db_path.resolve()
        self._lock = database_lock(self.db_path)

    def list_metrics(self) -> list[dict[str, str]]:
        with self._lock:
            connection = duckdb.connect(str(self.db_path), read_only=True)
            try:
                return [
                    {
                        "metric_id": str(metric_id),
                        "metric_name": str(name),
                        "description": str(description),
                        "unit": str(unit),
                    }
                    for metric_id, name, description, unit in connection.execute(
                        "SELECT metric_id, metric_name, description, unit "
                        "FROM metrics ORDER BY metric_id"
                    ).fetchall()
                ]
            finally:
                connection.close()

    def preview(self, content: bytes, file_name: str) -> dict[str, Any]:
        with self._lock:
            connection = duckdb.connect(str(self.db_path), read_only=True)
            try:
                catalog, organizations = self._catalog(connection)
                parsed = self._parse(content, file_name, catalog, organizations)
                preview = self._compare(connection, parsed)
                preview.pop("_existing", None)
                return preview
            finally:
                connection.close()

    def publish(
        self,
        content: bytes,
        file_name: str,
        *,
        confirm_overwrite: bool,
    ) -> dict[str, Any]:
        with self._lock:
            connection = duckdb.connect(str(self.db_path))
            try:
                catalog, organizations = self._catalog(connection)
                parsed = self._parse(content, file_name, catalog, organizations)
                preview = self._compare(connection, parsed)
                existing: dict[tuple[str, str, str], float] = preview.pop("_existing")
                if not preview["valid"]:
                    raise ValueError("文件校验失败：" + "；".join(preview["errors"][:3]))
                if preview["overwrite_count"] and not confirm_overwrite:
                    raise OverwriteConfirmationRequired(preview)

                rows: tuple[ImportRow, ...] = parsed["rows"]
                inserts = [row for row in rows if row.key not in existing]
                overwrites = [
                    row
                    for row in rows
                    if row.key in existing
                    and not _same_value(existing[row.key], row.metric_value)
                ]
                connection.execute("BEGIN TRANSACTION")
                try:
                    if inserts:
                        connection.executemany(
                            "INSERT INTO metric_values "
                            "(data_date, metric_id, org_id, metric_value) VALUES (?, ?, ?, ?)",
                            [
                                (row.data_date, row.metric_id, row.org_id, row.metric_value)
                                for row in inserts
                            ],
                        )
                    if overwrites:
                        connection.executemany(
                            "UPDATE metric_values SET metric_value = ? "
                            "WHERE data_date = ? AND metric_id = ? AND org_id = ?",
                            [
                                (row.metric_value, row.data_date, row.metric_id, row.org_id)
                                for row in overwrites
                            ],
                        )
                    connection.execute("COMMIT")
                except Exception:
                    connection.execute("ROLLBACK")
                    raise

                preview["published"] = True
                preview["overwrite_details"] = [
                    {
                        "data_date": row.data_date,
                        "metric_id": row.metric_id,
                        "org_id": row.org_id,
                        "old_value": existing[row.key],
                        "new_value": row.metric_value,
                    }
                    for row in overwrites[:100]
                ]
                return preview
            finally:
                connection.close()

    @staticmethod
    def _catalog(
        connection: duckdb.DuckDBPyConnection,
    ) -> tuple[dict[str, dict[str, str]], set[str]]:
        catalog = {
            str(metric_id): {
                "metric_id": str(metric_id),
                "metric_name": str(name),
                "description": str(description),
                "unit": str(unit),
            }
            for metric_id, name, description, unit in connection.execute(
                "SELECT metric_id, metric_name, description, unit FROM metrics"
            ).fetchall()
        }
        organizations = {
            str(row[0]) for row in connection.execute("SELECT org_id FROM organizations").fetchall()
        }
        return catalog, organizations

    def _parse(
        self,
        content: bytes,
        file_name: str,
        catalog: dict[str, dict[str, str]],
        organizations: set[str],
    ) -> dict[str, Any]:
        if not file_name.lower().endswith(".xlsx"):
            raise ValueError("仅支持.xlsx格式的Excel文件")
        if not content:
            raise ValueError("上传文件为空")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("无法读取Excel文件，请确认文件未损坏") from exc
        try:
            if len(workbook.worksheets) == 1:
                sheet = workbook.worksheets[0]
            elif DATA_SHEET in workbook.sheetnames:
                sheet = workbook[DATA_SHEET]
            else:
                raise ValueError(f"Excel中缺少“{DATA_SHEET}”工作表")

            errors: list[str] = []
            upload_mode = "complete_workbook" if METRIC_SHEET in workbook.sheetnames else "metric_data_only"
            uploaded_catalog: dict[str, str] | None = None
            if METRIC_SHEET in workbook.sheetnames:
                uploaded_catalog = self._parse_metric_sheet(workbook[METRIC_SHEET], catalog, errors)

            iterator = sheet.iter_rows(values_only=True)
            headers = tuple(_text(value) for value in (next(iterator, ()) or ()))
            if headers[: len(DATA_HEADERS)] != DATA_HEADERS:
                raise ValueError("指标数据表表头必须依次为：" + "、".join(DATA_HEADERS))

            unique_rows: dict[tuple[str, str, str], ImportRow] = {}
            duplicate_count = 0
            conflicting_duplicate_count = 0
            total_rows = 0
            for row_number, values in enumerate(iterator, start=2):
                cells = tuple(values[:5]) + (None,) * max(0, 5 - len(values))
                if all(value is None or _text(value) == "" for value in cells[:5]):
                    continue
                total_rows += 1
                raw_date, raw_metric_id, raw_metric_name, raw_org_id, raw_value = cells[:5]
                metric_id = _text(raw_metric_id)
                metric_name = _text(raw_metric_name)
                org_id = _text(raw_org_id)
                data_date = _date_value(raw_date)
                metric_value = _number_value(raw_value)

                missing = [
                    label
                    for label, value in (
                        ("数据日期", raw_date),
                        ("指标编号", metric_id),
                        ("指标名称", metric_name),
                        ("机构编号", org_id),
                        ("指标值", raw_value),
                    )
                    if value is None or _text(value) == ""
                ]
                if missing:
                    _append_error(errors, f"第{row_number}行关键字段为空：{'、'.join(missing)}")
                    continue
                if data_date is None:
                    _append_error(errors, f"第{row_number}行数据日期格式不合法")
                    continue
                if metric_id not in catalog:
                    _append_error(errors, f"第{row_number}行包含当前指标清单之外的指标：{metric_id}")
                    continue
                if metric_name != catalog[metric_id]["metric_name"]:
                    _append_error(
                        errors,
                        f"第{row_number}行指标名称与当前清单不一致：{metric_id}",
                    )
                    continue
                if uploaded_catalog is not None and metric_id not in uploaded_catalog:
                    _append_error(errors, f"第{row_number}行指标未包含在上传的指标清单表中：{metric_id}")
                    continue
                if org_id not in organizations:
                    _append_error(errors, f"第{row_number}行机构编号不存在：{org_id}")
                    continue
                if metric_value is None:
                    _append_error(errors, f"第{row_number}行指标值不是有效数值")
                    continue

                parsed_row = ImportRow(data_date, metric_id, metric_name, org_id, metric_value)
                previous = unique_rows.get(parsed_row.key)
                if previous is None:
                    unique_rows[parsed_row.key] = parsed_row
                elif _same_value(previous.metric_value, parsed_row.metric_value):
                    duplicate_count += 1
                else:
                    conflicting_duplicate_count += 1
                    _append_error(
                        errors,
                        f"第{row_number}行与文件内其他记录的日期、指标、机构相同，但指标值不同",
                    )

            if total_rows == 0:
                _append_error(errors, "指标数据表中没有可导入的数据")
            return {
                "file_name": file_name,
                "file_hash": hashlib.sha256(content).hexdigest(),
                "upload_mode": upload_mode,
                "total_rows": total_rows,
                "rows": tuple(unique_rows.values()),
                "duplicate_count": duplicate_count,
                "conflicting_duplicate_count": conflicting_duplicate_count,
                "errors": errors,
            }
        finally:
            workbook.close()

    @staticmethod
    def _parse_metric_sheet(
        sheet,
        catalog: dict[str, dict[str, str]],
        errors: list[str],
    ) -> dict[str, str]:
        iterator = sheet.iter_rows(values_only=True)
        headers = tuple(_text(value) for value in (next(iterator, ()) or ()))
        if headers[: len(METRIC_HEADERS)] != METRIC_HEADERS:
            raise ValueError("指标清单表表头必须依次为：" + "、".join(METRIC_HEADERS))
        uploaded: dict[str, str] = {}
        for row_number, values in enumerate(iterator, start=2):
            metric_id = _text(values[0] if len(values) > 0 else None)
            metric_name = _text(values[1] if len(values) > 1 else None)
            if not metric_id and not metric_name:
                continue
            if not metric_id or not metric_name:
                _append_error(errors, f"指标清单表第{row_number}行指标编号或名称为空")
                continue
            if metric_id not in catalog:
                _append_error(errors, f"指标清单表包含当前不存在的指标：{metric_id}")
                continue
            if metric_name != catalog[metric_id]["metric_name"]:
                _append_error(errors, f"指标清单表中的指标名称与当前清单不一致：{metric_id}")
                continue
            previous = uploaded.get(metric_id)
            if previous is not None and previous != metric_name:
                _append_error(errors, f"指标清单表中的指标定义重复且不一致：{metric_id}")
                continue
            uploaded[metric_id] = metric_name
        if not uploaded:
            _append_error(errors, "指标清单表中没有有效指标")
        return uploaded

    @staticmethod
    def _existing_values(
        connection: duckdb.DuckDBPyConnection,
        rows: tuple[ImportRow, ...],
    ) -> dict[tuple[str, str, str], float]:
        if not rows:
            return {}
        metrics = sorted({row.metric_id for row in rows})
        dates = sorted({row.data_date for row in rows})
        placeholders = ", ".join("?" for _ in metrics)
        result = connection.execute(
            "SELECT data_date, metric_id, org_id, metric_value FROM metric_values "
            f"WHERE data_date BETWEEN ? AND ? AND metric_id IN ({placeholders})",
            [dates[0], dates[-1], *metrics],
        ).fetchall()
        return {
            (str(data_date), str(metric_id), str(org_id)): float(metric_value)
            for data_date, metric_id, org_id, metric_value in result
        }

    def _compare(
        self,
        connection: duckdb.DuckDBPyConnection,
        parsed: dict[str, Any],
    ) -> dict[str, Any]:
        rows: tuple[ImportRow, ...] = parsed["rows"]
        existing = self._existing_values(connection, rows)
        insert_count = 0
        unchanged_count = 0
        overwrite_count = 0
        overwrite_samples: list[dict[str, object]] = []
        for row in rows:
            current = existing.get(row.key)
            if current is None:
                insert_count += 1
            elif _same_value(current, row.metric_value):
                unchanged_count += 1
            else:
                overwrite_count += 1
                if len(overwrite_samples) < 20:
                    overwrite_samples.append(
                        {
                            "data_date": row.data_date,
                            "metric_id": row.metric_id,
                            "org_id": row.org_id,
                            "old_value": current,
                            "new_value": row.metric_value,
                        }
                    )
        touched_metrics = sorted({row.metric_id for row in rows})
        touched_orgs = {row.org_id for row in rows}
        dates = sorted({row.data_date for row in rows})
        errors = list(parsed["errors"])
        return {
            "file_name": parsed["file_name"],
            "file_hash": parsed["file_hash"],
            "upload_mode": parsed["upload_mode"],
            "valid": not errors,
            "can_publish": not errors and bool(rows),
            "total_rows": parsed["total_rows"],
            "parsed_row_count": len(rows),
            "insert_count": insert_count,
            "unchanged_count": unchanged_count,
            "overwrite_count": overwrite_count,
            "duplicate_count": parsed["duplicate_count"],
            "conflicting_duplicate_count": parsed["conflicting_duplicate_count"],
            "metric_ids": touched_metrics,
            "organization_count": len(touched_orgs),
            "date_start": dates[0] if dates else None,
            "date_end": dates[-1] if dates else None,
            "errors": errors,
            "overwrite_samples": overwrite_samples,
            "_existing": existing,
        }
