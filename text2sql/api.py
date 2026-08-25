"""FastAPI transport for the Text2SQL service."""

from __future__ import annotations

from dataclasses import replace
from functools import lru_cache

from io import BytesIO, StringIO
import csv
import json
import os
import re
from typing import Any
import uuid

from fastapi import FastAPI, Header, HTTPException, Query, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.openapi.utils import get_openapi
from fastapi.responses import JSONResponse, Response, StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from .errors import Text2SQLError
from .product_auth import ProductAuthService
from .product_models import (
    BusinessRole,
    ProductQueryResponse,
    TimeGranularity,
    UserContext,
)
from .product_service import (
    AccountFrozenError,
    LargeExportConfirmationRequired,
    ProductQueryService,
)
from .product_store import ProductStore
from .service import Text2SQLService


API_TAGS = [
    {"name": "系统", "description": "健康检查与能力发现"},
    {"name": "认证", "description": "本地演示认证；生产环境应接入银行统一认证"},
    {"name": "智能问数", "description": "同步或流式执行自然语言数据查询"},
    {"name": "会话", "description": "会话状态和标题管理"},
    {"name": "历史", "description": "查询历史、详情与批量导出"},
    {"name": "分享", "description": "限时分享查询结果"},
    {"name": "管理", "description": "管理概览和操作审计"},
    {"name": "兼容接口", "description": "旧版接口，仅用于兼容现有调用方"},
]
app = FastAPI(
    title="Bank Text2SQL",
    version="0.1.0",
    description="银行智能问数标准化 API。生产环境请通过 API 网关访问。",
    openapi_tags=API_TAGS,
)
cors_origins = [
    value.strip()
    for value in os.getenv(
        "TEXT2SQL_CORS_ORIGINS",
        (
            "http://localhost:3000,http://127.0.0.1:3000,"
            "http://localhost:3101,http://127.0.0.1:3101"
        ),
    ).split(",")
    if value.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


REQUEST_ID_PATTERN = re.compile(r"^[A-Za-z0-9._:-]{1,128}$")
ERROR_CODE_BY_STATUS = {
    400: "REQUEST_INVALID",
    401: "AUTH_UNAUTHORIZED",
    403: "AUTH_FORBIDDEN",
    409: "CONFIRMATION_REQUIRED",
    410: "LEGACY_ENDPOINT_DISABLED",
    423: "ACCOUNT_FROZEN",
    404: "RESOURCE_NOT_FOUND",
    422: "REQUEST_VALIDATION_ERROR",
    429: "RATE_LIMITED",
    500: "SYSTEM_ERROR",
    504: "QUERY_TIMEOUT",
}


class ErrorResponse(BaseModel):
    code: str
    detail: Any
    request_id: str


ERROR_RESPONSES = {
    status: {"model": ErrorResponse}
    for status in (400, 401, 403, 404, 409, 410, 422, 423, 500)
}
SSE_RESPONSES = {
    200: {
        "description": "Server-Sent Events stream",
        "content": {"text/event-stream": {"schema": {"type": "string"}}},
    },
    **ERROR_RESPONSES,
}


def _new_request_id(candidate: str | None = None) -> str:
    if candidate and REQUEST_ID_PATTERN.fullmatch(candidate):
        return candidate
    return uuid.uuid4().hex


def _request_id(request: Request, preferred: str | None = None) -> str:
    value = _new_request_id(preferred or getattr(request.state, "request_id", None))
    request.state.request_id = value
    return value


@app.middleware("http")
async def request_id_middleware(request: Request, call_next):
    request.state.request_id = _new_request_id(request.headers.get("X-Request-ID"))
    response = await call_next(request)
    response.headers["X-Request-ID"] = request.state.request_id
    return response


@app.exception_handler(HTTPException)
async def http_error_response(request: Request, exc: HTTPException) -> JSONResponse:
    return JSONResponse(
        status_code=exc.status_code,
        content={
            "code": ERROR_CODE_BY_STATUS.get(exc.status_code, "REQUEST_FAILED"),
            "detail": exc.detail,
            "request_id": _request_id(request),
        },
        headers=exc.headers,
    )


@app.exception_handler(RequestValidationError)
async def validation_error_response(
    request: Request,
    exc: RequestValidationError,
) -> JSONResponse:
    return JSONResponse(
        status_code=422,
        content={
            "code": ERROR_CODE_BY_STATUS[422],
            "detail": exc.errors(),
            "request_id": _request_id(request),
        },
    )


class QueryRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    session_id: str | None = Field(default=None, max_length=128)
    request_id: str | None = Field(
        default=None,
        min_length=1,
        max_length=128,
        pattern=r"^[A-Za-z0-9._:-]+$",
    )
    caller_system: str | None = Field(
        default=None,
        min_length=1,
        max_length=64,
        pattern=r"^[A-Za-z0-9._:-]+$",
    )


class DrillRequest(BaseModel):
    target_granularity: TimeGranularity
    period_start: str | None = None
    period_end: str | None = None


class ShareRequest(BaseModel):
    expires_hours: int = Field(default=24, ge=1, le=168)


class LoginRequest(BaseModel):
    username: str = Field(min_length=1, max_length=128)
    password: str = Field(min_length=1, max_length=256)


class RenameSessionRequest(BaseModel):
    title: str = Field(min_length=1, max_length=60)
    owner_user_id: str | None = Field(default=None, max_length=128)


class AlertStatusRequest(BaseModel):
    status: str = Field(pattern="^(acknowledged|resolved)$")


@lru_cache(maxsize=1)
def get_service() -> Text2SQLService:
    return Text2SQLService()


@lru_cache(maxsize=1)
def get_product_service() -> ProductQueryService:
    service = get_service()
    return ProductQueryService(service, ProductStore(service.settings.product_db_path))


@lru_cache(maxsize=1)
def get_auth_service() -> ProductAuthService:
    return ProductAuthService()


def user_context(
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = None,
) -> UserContext:
    if authorization:
        try:
            return get_auth_service().resolve_bearer(authorization).to_context()
        except PermissionError as exc:
            raise HTTPException(status_code=401, detail=str(exc)) from exc
    allow_header_auth = os.getenv("TEXT2SQL_ALLOW_HEADER_AUTH", "false").strip().lower()
    if allow_header_auth not in {"1", "true", "yes", "on"}:
        raise HTTPException(status_code=401, detail="请先登录")
    business_roles: set[BusinessRole] = {
        "head_office_manager",
        "branch_manager",
        "business_staff",
        "risk_compliance",
        "finance_staff",
        "system_admin",
    }
    if x_user_role not in {"viewer", "analyst", "admin", *business_roles}:
        raise HTTPException(status_code=400, detail="X-User-Role不是受支持的角色")
    organizations = tuple(
        value.strip() for value in x_org_scope.split(",") if value.strip()
    )
    business_role: BusinessRole | None
    if x_user_role in business_roles:
        business_role = x_user_role  # type: ignore[assignment]
        role = "admin" if business_role == "system_admin" else "analyst"
    else:
        role = x_user_role
        business_role = (
            "branch_manager"
            if role == "analyst" and organizations
            else "head_office_manager"
            if role == "analyst"
            else "system_admin"
            if role == "admin"
            else None
        )
    organization_access = (
        "restricted"
        if business_role in {"branch_manager", "business_staff"}
        else "none"
        if business_role == "system_admin"
        else "all"
    )
    return UserContext(
        user_id=x_user_id,
        role=role,
        allowed_organizations=organizations,
        business_role=business_role,
        organization_access=organization_access,
    )


@app.get("/health", tags=["系统"], summary="健康检查")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post(
    "/api/v1/auth/login",
    tags=["认证"],
    summary="登录并获取访问令牌",
    responses=ERROR_RESPONSES,
)
def login(request: LoginRequest) -> dict[str, object]:
    try:
        token, user = get_auth_service().login(request.username, request.password)
    except ValueError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    return {"access_token": token, "token_type": "bearer", "user": user.to_dict()}


@app.get(
    "/api/v1/auth/me",
    tags=["认证"],
    summary="获取当前用户",
    responses=ERROR_RESPONSES,
)
def current_user(
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    try:
        return get_auth_service().resolve_bearer(authorization).to_dict()
    except PermissionError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc


@app.post(
    "/api/v1/auth/logout",
    tags=["认证"],
    summary="注销当前令牌",
    responses=ERROR_RESPONSES,
)
def logout(authorization: str | None = Header(default=None)) -> dict[str, bool]:
    try:
        get_auth_service().logout(authorization)
    except PermissionError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    return {"logged_out": True}


@app.post("/query", tags=["兼容接口"], summary="旧版智能问数")
def query(request: QueryRequest) -> dict[str, object]:
    raise HTTPException(
        status_code=410,
        detail="旧版/query已关闭，请登录后使用/api/v1/queries",
    )


@app.delete(
    "/sessions/{session_id}",
    tags=["兼容接口"],
    summary="清除旧版会话",
)
def clear_session(session_id: str) -> dict[str, bool]:
    raise HTTPException(
        status_code=410,
        detail="旧版会话接口已关闭，请使用/api/v1/sessions",
    )


def _product_error(exc: Exception) -> HTTPException:
    if isinstance(exc, AccountFrozenError):
        return HTTPException(
            status_code=423,
            detail={
                "message": str(exc),
                "frozen_until": exc.frozen_until,
                "reason": exc.reason,
            },
        )
    if isinstance(exc, LargeExportConfirmationRequired):
        return HTTPException(
            status_code=409,
            detail={
                "message": str(exc),
                "row_count": exc.row_count,
                "confirm_parameter": "confirm_large_export=true",
            },
        )
    if isinstance(exc, PermissionError):
        return HTTPException(status_code=403, detail=str(exc))
    if isinstance(exc, (Text2SQLError, ValueError)):
        return HTTPException(status_code=400, detail=str(exc))
    return HTTPException(status_code=500, detail="系统处理失败")


@app.get(
    "/api/v1/capabilities",
    tags=["系统"],
    summary="获取服务能力",
)
def capabilities() -> dict[str, object]:
    return {
        "api_version": "v1",
        "features": {
            "nl2sql": True,
            "chart_recommendation": True,
            "business_insight": True,
            "final_answer_stream": True,
            "time_drill": ["year", "quarter", "month", "day"],
            "history": True,
            "export": ["xlsx", "csv"],
            "share": True,
            "audit": True,
            "fine_grained_permissions": True,
            "dynamic_masking": True,
            "security_alerts": True,
        },
        "time_aggregation": "period_end",
    }


@app.delete(
    "/api/v1/sessions/{session_id}",
    tags=["会话"],
    summary="清除会话",
    responses=ERROR_RESPONSES,
)
def clear_product_session(
    session_id: str,
    delete_history: bool = Query(default=True),
    owner_user_id: str | None = Query(default=None, max_length=128),
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        get_product_service().clear_session(
            session_id,
            user,
            delete_history=delete_history,
            owner_user_id=owner_user_id,
        )
        return {"cleared": True}
    except Exception as exc:
        raise _product_error(exc) from exc


@app.patch(
    "/api/v1/sessions/{session_id}",
    tags=["会话"],
    summary="重命名会话",
    responses=ERROR_RESPONSES,
)
def rename_product_session(
    session_id: str,
    request: RenameSessionRequest,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        renamed = get_product_service().rename_session(
            session_id,
            request.title,
            user,
            owner_user_id=request.owner_user_id,
        )
        return {"renamed": True, **renamed}
    except Exception as exc:
        raise _product_error(exc) from exc


@app.post(
    "/api/v1/queries",
    response_model=ProductQueryResponse,
    tags=["智能问数"],
    summary="执行同步自然语言查询",
    responses=ERROR_RESPONSES,
)
def product_query(
    request: QueryRequest,
    http_request: Request,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    integration_request_id = _request_id(http_request, request.request_id)
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        response = get_product_service().query(
            request.question,
            request.session_id,
            user,
        )
        return replace(
            response,
            request_id=integration_request_id,
            caller_system=request.caller_system,
        ).to_dict()
    except Exception as exc:
        raise _product_error(exc) from exc


@app.post(
    "/api/v1/queries/stream",
    tags=["智能问数"],
    summary="流式执行自然语言查询",
    responses=SSE_RESPONSES,
)
def product_query_stream(
    request: QueryRequest,
    http_request: Request,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> StreamingResponse:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    public_session_id = request.session_id or uuid.uuid4().hex
    integration_request_id = _request_id(http_request, request.request_id)

    def event_stream():
        yield _sse(
            "status",
            {
                "stage": "understanding",
                "message": "正在理解问题",
                "session_id": public_session_id,
                "request_id": integration_request_id,
            },
        )
        try:
            result = get_product_service().query(
                request.question,
                public_session_id,
                user,
            ).to_dict()
            result["request_id"] = integration_request_id
            result["caller_system"] = request.caller_system
            yield _sse("status", {"stage": "data_ready", "message": "数据查询完成"})
            yield _sse("result", result)
            yield _sse(
                "complete",
                {
                    "query_id": result["query_id"],
                    "answer_mode": result["answer_mode"],
                    "answer_status": result["answer_status"],
                    "request_id": integration_request_id,
                },
            )
        except Exception as exc:
            status = (
                423
                if isinstance(exc, AccountFrozenError)
                else 403
                if isinstance(exc, PermissionError)
                else 400
            )
            message = (
                str(exc)
                if isinstance(exc, (PermissionError, Text2SQLError, ValueError))
                else "系统处理失败"
            )
            yield _sse(
                "error",
                {
                    "status": status,
                    "message": message,
                    "session_id": public_session_id,
                    "code": ERROR_CODE_BY_STATUS.get(status, "REQUEST_FAILED"),
                    "request_id": integration_request_id,
                },
            )

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


@app.get(
    "/api/v1/queries/{query_id}/answer/stream",
    tags=["智能问数"],
    summary="流式生成最终回答",
    responses=SSE_RESPONSES,
)
def final_answer_stream(
    query_id: str,
    http_request: Request,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> StreamingResponse:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)

    def event_stream():
        try:
            for event, payload in get_product_service().stream_final_answer(
                query_id,
                user,
            ):
                yield _sse(event, payload)
        except Exception as exc:
            status = (
                423
                if isinstance(exc, AccountFrozenError)
                else 403
                if isinstance(exc, PermissionError)
                else 400
            )
            message = (
                str(exc)
                if isinstance(exc, (PermissionError, Text2SQLError, ValueError))
                else "最终回答生成失败"
            )
            yield _sse(
                "answer_error",
                {
                    "query_id": query_id,
                    "status": status,
                    "message": message,
                    "code": ERROR_CODE_BY_STATUS.get(status, "REQUEST_FAILED"),
                    "request_id": _request_id(http_request),
                },
            )

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


@app.post(
    "/api/v1/queries/{query_id}/drill",
    response_model=ProductQueryResponse,
    tags=["智能问数"],
    summary="执行时间下钻",
    responses=ERROR_RESPONSES,
)
def time_drill(
    query_id: str,
    request: DrillRequest,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return get_product_service().time_drill(
            query_id,
            request.target_granularity,
            user,
            period_start=request.period_start,
            period_end=request.period_end,
        ).to_dict()
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/history",
    tags=["历史"],
    summary="查询历史列表",
    responses=ERROR_RESPONSES,
)
def history(
    limit: int | None = Query(default=None, ge=1, le=5000),
    keyword: str | None = Query(default=None, max_length=200),
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return {
            "items": get_product_service().history(user, limit=limit, keyword=keyword)
        }
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/history/{query_id}",
    response_model=ProductQueryResponse,
    tags=["历史"],
    summary="查询历史详情",
    responses=ERROR_RESPONSES,
)
def history_detail(
    query_id: str,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return get_product_service().get_query(query_id, user).to_dict()
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/sessions/{session_id}/history",
    tags=["历史"],
    summary="查询会话历史",
    responses=ERROR_RESPONSES,
)
def session_history(
    session_id: str,
    owner_user_id: str | None = Query(default=None, max_length=128),
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return {
            "items": [
                item.to_dict()
                for item in get_product_service().session_history(
                    session_id,
                    user,
                    owner_user_id=owner_user_id,
                )
            ]
        }
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/history/export/batch",
    tags=["历史"],
    summary="批量导出历史记录",
    responses=ERROR_RESPONSES,
)
def batch_export_history(
    scope: str = Query(pattern="^(all|recent|session)$"),
    recent_count: int = Query(default=20, ge=1, le=5000),
    session_id: str | None = Query(default=None, max_length=128),
    owner_user_id: str | None = Query(default=None, max_length=128),
    confirm_large_export: bool = Query(default=False),
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> Response:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        records = get_product_service().batch_export_records(
            user,
            scope=scope,
            recent_count=recent_count,
            session_id=session_id,
            owner_user_id=owner_user_id,
            confirm_large_export=confirm_large_export,
        )
    except Exception as exc:
        raise _product_error(exc) from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "历史汇总"
    summary.append(
        [
            "查询编号",
            "用户",
            "thread_id",
            "生成时间",
            "问题",
            "最终回答",
            "路径",
            "结果行数",
            "耗时(ms)",
        ]
    )
    details = workbook.create_sheet("完整查询数据")
    details.append(
        [
            "查询编号",
            "用户",
            "thread_id",
            "问题",
            "SQL",
            "查询计划JSON",
            "字段JSON",
            "结果JSON",
            "图表配置JSON",
            "洞察JSON",
            "告警JSON",
        ]
    )
    for record in records:
        summary.append(
            [
                record["query_id"],
                record["user_id"],
                record["session_id"],
                record["created_at"],
                record["question"],
                record["answer"],
                record["route"],
                len(record["rows"]),
                record["duration_ms"],
            ]
        )
        details.append(
            [
                record["query_id"],
                record["user_id"],
                record["session_id"],
                record["question"],
                record.get("sql") or "",
                json.dumps(record["plan"], ensure_ascii=False, default=str),
                json.dumps(record["columns"], ensure_ascii=False, default=str),
                json.dumps(record["rows"], ensure_ascii=False, default=str),
                json.dumps(record["visualization"], ensure_ascii=False, default=str),
                json.dumps(record["insight"], ensure_ascii=False, default=str),
                json.dumps(record["warnings"], ensure_ascii=False, default=str),
            ]
        )
    summary.freeze_panes = "A2"
    details.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="bank-history-batch.xlsx"'
        },
    )


@app.get(
    "/api/v1/queries/{query_id}/export",
    tags=["历史"],
    summary="导出单次查询",
    responses=ERROR_RESPONSES,
)
def export_query(
    query_id: str,
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    confirm_large_export: bool = Query(default=False),
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> Response:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        record = get_product_service().export_record(
            query_id,
            user,
            confirm_large_export=confirm_large_export,
        )
    except Exception as exc:
        raise _product_error(exc) from exc
    filename = f"bank-query-{query_id[:8]}"
    if format == "csv":
        stream = StringIO()
        writer = csv.writer(stream)
        writer.writerow(record["columns"])
        writer.writerows(record["rows"])
        return Response(
            content="\ufeff" + stream.getvalue(),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'},
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "查询结果"
    sheet.append(["问题", record["question"]])
    sheet.append(["结论", record["answer"]])
    sheet.append(["生成时间", record["created_at"]])
    sheet.append([])
    sheet.append(record["columns"])
    for row in record["rows"]:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


@app.post(
    "/api/v1/queries/{query_id}/share",
    tags=["分享"],
    summary="创建限时分享",
    responses=ERROR_RESPONSES,
)
def share_query(
    query_id: str,
    request: ShareRequest,
    x_user_id: str = Header(default="demo-analyst", max_length=128),
    x_user_role: str = Header(default="analyst"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, str]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return get_product_service().create_share(
            query_id,
            user,
            expires_hours=request.expires_hours,
        )
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/shares/{token}",
    response_model=ProductQueryResponse,
    tags=["分享"],
    summary="访问分享结果",
    responses=ERROR_RESPONSES,
)
def shared_query(
    token: str,
    x_user_id: str = Header(default="demo-viewer", max_length=128),
    x_user_role: str = Header(default="viewer"),
    x_org_scope: str = Header(default="", max_length=2000),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, x_org_scope, authorization)
    try:
        return get_product_service().resolve_share(token, user).to_dict()
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/admin/overview",
    tags=["管理"],
    summary="管理概览",
    responses=ERROR_RESPONSES,
)
def admin_overview(
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return get_product_service().admin_overview(user)
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/admin/audit",
    tags=["管理"],
    summary="查询操作审计",
    responses=ERROR_RESPONSES,
)
def admin_audit(
    sort: str = Query(default="newest", pattern="^(newest|risk_desc|risk_asc)$"),
    risk: str = Query(default="all", pattern="^(all|elevated|low|medium|high)$"),
    audit_user: str | None = Query(default=None, max_length=128, alias="user"),
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    current_user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return {
            "items": get_product_service().admin_audit(
                current_user,
                sort_by=sort,
                risk_level=risk,
                user_id=audit_user,
            )
        }
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/admin/alerts",
    tags=["管理"],
    summary="查询安全告警",
    responses=ERROR_RESPONSES,
)
def admin_alerts(
    status: str = Query(default="all", pattern="^(all|open|acknowledged|resolved)$"),
    severity: str = Query(default="all", pattern="^(all|low|medium|high)$"),
    alert_user: str | None = Query(default=None, max_length=128, alias="user"),
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    current_user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return {
            "items": get_product_service().admin_alerts(
                current_user,
                status=status,
                severity=severity,
                user_id=alert_user,
            )
        }
    except Exception as exc:
        raise _product_error(exc) from exc


@app.patch(
    "/api/v1/admin/alerts/{alert_id}",
    tags=["管理"],
    summary="处置安全告警",
    responses=ERROR_RESPONSES,
)
def update_admin_alert(
    alert_id: str,
    request: AlertStatusRequest,
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, str]:
    current_user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return get_product_service().update_alert_status(
            alert_id,
            request.status,
            current_user,
        )
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/admin/freezes",
    tags=["管理"],
    summary="查询当前冻结账号",
    responses=ERROR_RESPONSES,
)
def admin_freezes(
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return {"items": get_product_service().admin_freezes(user)}
    except Exception as exc:
        raise _product_error(exc) from exc


@app.delete(
    "/api/v1/admin/freezes/{user_id}",
    tags=["管理"],
    summary="手动解冻账号",
    responses=ERROR_RESPONSES,
)
def unfreeze_admin_user(
    user_id: str,
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    current_user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return get_product_service().unfreeze_user(user_id, current_user)
    except Exception as exc:
        raise _product_error(exc) from exc


@app.get(
    "/api/v1/admin/users",
    tags=["管理"],
    summary="查询审计用户",
    responses=ERROR_RESPONSES,
)
def admin_users(
    x_user_id: str = Header(default="demo-admin", max_length=128),
    x_user_role: str = Header(default="admin"),
    authorization: str | None = Header(default=None),
) -> dict[str, object]:
    user = user_context(x_user_id, x_user_role, "", authorization)
    try:
        return {"items": get_product_service().admin_users(user)}
    except Exception as exc:
        raise _product_error(exc) from exc


def custom_openapi() -> dict[str, object]:
    if app.openapi_schema:
        return app.openapi_schema
    schema = get_openapi(
        title=app.title,
        version=app.version,
        description=app.description,
        routes=app.routes,
        tags=app.openapi_tags,
    )
    components = schema.setdefault("components", {})
    security_schemes = components.setdefault("securitySchemes", {})
    security_schemes["BearerAuth"] = {
        "type": "http",
        "scheme": "bearer",
        "bearerFormat": "opaque-token",
    }
    public_operations = {
        ("/api/v1/auth/login", "post"),
        ("/api/v1/capabilities", "get"),
        ("/api/v1/shares/{token}", "get"),
    }
    for path, path_item in schema.get("paths", {}).items():
        if not path.startswith("/api/v1"):
            continue
        for method, operation in path_item.items():
            if method not in {"get", "post", "put", "patch", "delete"}:
                continue
            if (path, method) not in public_operations:
                operation["security"] = [{"BearerAuth": []}]
    app.openapi_schema = schema
    return schema


app.openapi = custom_openapi


def _sse(event: str, data: object) -> str:
    return (
        f"event: {event}\n"
        f"data: {json.dumps(data, ensure_ascii=False, default=str)}\n\n"
    )
