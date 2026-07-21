"""Build the read-only analytical DuckDB from the competition workbook."""

from __future__ import annotations

from datetime import date, datetime
import csv
import os
from pathlib import Path
import tempfile

import duckdb
from openpyxl import load_workbook

from .errors import DataBuildError


SCHEMA_VERSION = "1"


def _cell_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _sheet_rows(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    next(iterator, None)
    yield from iterator


def _database_is_current(db_path: Path, xlsx_path: Path) -> bool:
    if not db_path.exists():
        return False
    try:
        connection = duckdb.connect(str(db_path), read_only=True)
        row = connection.execute(
            "SELECT source_size, source_mtime_ns, schema_version FROM build_metadata LIMIT 1"
        ).fetchone()
        connection.close()
        stat = xlsx_path.stat()
        return bool(
            row
            and int(row[0]) == stat.st_size
            and int(row[1]) == stat.st_mtime_ns
            and str(row[2]) == SCHEMA_VERSION
        )
    except Exception:
        return False


def ensure_database(xlsx_path: Path, db_path: Path, force: bool = False) -> Path:
    """Create a normalized database when it is missing or older than the workbook."""
    xlsx_path = xlsx_path.resolve()
    db_path = db_path.resolve()
    if not xlsx_path.exists():
        raise DataBuildError(f"数据集不存在：{xlsx_path}")
    if not force and _database_is_current(db_path, xlsx_path):
        return db_path
    db_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = db_path.with_suffix(db_path.suffix + ".building")
    temp_wal_path = Path(str(temp_path) + ".wal")
    if temp_path.exists():
        temp_path.unlink()
    if temp_wal_path.exists():
        temp_wal_path.unlink()
    workbook = None
    connection = None
    try:
        workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
        connection = duckdb.connect(str(temp_path))
        connection.execute("SET memory_limit='1GB'")
        connection.execute(
            """
            CREATE TABLE organizations (
                org_id VARCHAR PRIMARY KEY,
                org_name VARCHAR NOT NULL
            );
            CREATE TABLE metrics (
                metric_id VARCHAR PRIMARY KEY,
                metric_name VARCHAR NOT NULL,
                description VARCHAR NOT NULL,
                unit VARCHAR NOT NULL
            );
            CREATE TABLE derived_rules (
                rule_name VARCHAR PRIMARY KEY,
                description VARCHAR NOT NULL
            );
            CREATE TABLE metric_values (
                data_date DATE NOT NULL,
                metric_id VARCHAR NOT NULL,
                org_id VARCHAR NOT NULL,
                metric_value DOUBLE NOT NULL,
                PRIMARY KEY (data_date, metric_id, org_id)
            );
            CREATE TABLE build_metadata (
                source_path VARCHAR,
                source_size BIGINT,
                source_mtime_ns BIGINT,
                schema_version VARCHAR,
                built_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            """
        )
        organizations = [
            (str(row[0]).strip(), str(row[1]).strip())
            for row in _sheet_rows(workbook, "机构信息表")
            if row[0] and row[1]
        ]
        metrics = [
            (str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip(), str(row[3]).strip())
            for row in _sheet_rows(workbook, "指标清单表")
            if row[0] and row[1]
        ]
        derived_rules = [
            (str(row[0]).strip(), str(row[1]).strip())
            for row in _sheet_rows(workbook, "衍生维度说明")
            if row[0] and row[1]
        ]
        connection.executemany("INSERT INTO organizations VALUES (?, ?)", organizations)
        connection.executemany("INSERT INTO metrics VALUES (?, ?, ?, ?)", metrics)
        connection.executemany("INSERT INTO derived_rules VALUES (?, ?)", derived_rules)

        with tempfile.TemporaryDirectory(prefix="text2sql_build_", dir=db_path.parent) as csv_dir:
            csv_path = Path(csv_dir) / "metric_values.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["data_date", "metric_id", "org_id", "metric_value"])
                for row in _sheet_rows(workbook, "指标数据表"):
                    if not row[0] or not row[1] or not row[3] or row[4] is None:
                        continue
                    writer.writerow(
                        [_cell_date(row[0]), str(row[1]).strip(), str(row[3]).strip(), float(row[4])]
                    )
            escaped_csv_path = str(csv_path).replace("'", "''")
            connection.execute(
                f"COPY metric_values FROM '{escaped_csv_path}' (HEADER, DELIMITER ',', QUOTE '\"')"
            )
        connection.execute("CREATE INDEX metric_lookup ON metric_values(metric_id, data_date, org_id)")
        stat = xlsx_path.stat()
        connection.execute(
            "INSERT INTO build_metadata(source_path, source_size, source_mtime_ns, schema_version) VALUES (?, ?, ?, ?)",
            [str(xlsx_path), stat.st_size, stat.st_mtime_ns, SCHEMA_VERSION],
        )
        connection.close()
        connection = None
        if db_path.exists():
            db_path.unlink()
        os.replace(temp_path, db_path)
        return db_path
    except Exception as exc:
        raise DataBuildError(f"构建 DuckDB 失败：{type(exc).__name__}: {exc}") from exc
    finally:
        if connection is not None:
            connection.close()
        if workbook is not None:
            workbook.close()
        if temp_path.exists():
            temp_path.unlink()
        if temp_wal_path.exists():
            temp_wal_path.unlink()
